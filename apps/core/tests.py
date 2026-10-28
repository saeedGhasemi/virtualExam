import io
import json
import tempfile
import uuid
from datetime import timedelta

from django.contrib.auth import get_user_model
from django.db import connection
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone


def _insert(table, columns, values):
    placeholders = ', '.join(['%s'] * len(values))
    column_list = ', '.join(columns)
    with connection.cursor() as cursor:
        cursor.execute(f'INSERT INTO {table} ({column_list}) VALUES ({placeholders})', values)


class _ScopedOrgFixture(TestCase):
    """فیکسچر مشترک: یک مدیر آموزشی با scope روی یک واحد سازمانی، به‌همراه استاد/دانشجوی داخل و خارج آن محدوده."""

    @classmethod
    def setUpTestData(cls):
        cls.ou_scoped = str(uuid.uuid4())
        cls.ou_other = str(uuid.uuid4())
        _insert('org_units', ['id', 'name', 'type'], [cls.ou_scoped, 'واحد تحت مدیریت', 'department'])
        _insert('org_units', ['id', 'name', 'type'], [cls.ou_other, 'واحد دیگر', 'department'])

        cls.manager_id = cls._make_profile('manager_test', 'مدیر آموزشی تست')
        _insert('academic_manager_profiles', ['user_id'], [cls.manager_id])
        _insert('academic_manager_scopes', ['id', 'manager_id', 'org_unit_id'],
                [str(uuid.uuid4()), cls.manager_id, cls.ou_scoped])
        _insert('user_roles', ['id', 'user_id', 'role'], [str(uuid.uuid4()), cls.manager_id, 'academic_manager'])

        cls.teacher_in_id = cls._make_profile('teacher_in_scope', 'استاد داخل محدوده')
        _insert('teacher_profiles', ['user_id', 'org_unit_id'], [cls.teacher_in_id, cls.ou_scoped])
        _insert('user_roles', ['id', 'user_id', 'role'], [str(uuid.uuid4()), cls.teacher_in_id, 'teacher'])

        cls.teacher_out_id = cls._make_profile('teacher_out_scope', 'استاد خارج محدوده')
        _insert('teacher_profiles', ['user_id', 'org_unit_id'], [cls.teacher_out_id, cls.ou_other])
        _insert('user_roles', ['id', 'user_id', 'role'], [str(uuid.uuid4()), cls.teacher_out_id, 'teacher'])

        cls.student_in_id = cls._make_profile('student_in_scope', 'دانشجوی داخل محدوده')
        _insert('student_profiles', ['user_id', 'org_unit_id'], [cls.student_in_id, cls.ou_scoped])
        _insert('user_roles', ['id', 'user_id', 'role'], [str(uuid.uuid4()), cls.student_in_id, 'student'])

        cls.student_out_id = cls._make_profile('student_out_scope', 'دانشجوی خارج محدوده')
        _insert('student_profiles', ['user_id', 'org_unit_id'], [cls.student_out_id, cls.ou_other])
        _insert('user_roles', ['id', 'user_id', 'role'], [str(uuid.uuid4()), cls.student_out_id, 'student'])

        cls.course_in_id = str(uuid.uuid4())
        _insert('courses', ['id', 'title', 'org_unit_id'], [cls.course_in_id, 'درس داخل محدوده', cls.ou_scoped])
        cls.course_out_id = str(uuid.uuid4())
        _insert('courses', ['id', 'title', 'org_unit_id'], [cls.course_out_id, 'درس خارج محدوده', cls.ou_other])

        cls.group_in_id = str(uuid.uuid4())
        _insert(
            'student_groups',
            ['id', 'teacher_id', 'course_id', 'course_name', 'academic_year'],
            [cls.group_in_id, cls.teacher_in_id, cls.course_in_id, 'درس داخل محدوده', '1403'],
        )
        cls.group_out_id = str(uuid.uuid4())
        _insert(
            'student_groups',
            ['id', 'teacher_id', 'course_id', 'course_name', 'academic_year'],
            [cls.group_out_id, cls.teacher_out_id, cls.course_out_id, 'درس خارج محدوده', '1403'],
        )

        now = timezone.now()
        cls.exam_in_id = str(uuid.uuid4())
        _insert('exams', ['id', 'title', 'teacher_id', 'course_id', 'start_at'], [cls.exam_in_id, 'آزمون داخل محدوده', cls.teacher_in_id, cls.course_in_id, now])
        cls.exam_out_id = str(uuid.uuid4())
        _insert('exams', ['id', 'title', 'teacher_id', 'course_id', 'start_at'], [cls.exam_out_id, 'آزمون خارج محدوده', cls.teacher_out_id, cls.course_out_id, now])

        cls.question_in_id = str(uuid.uuid4())
        _insert('questions', ['id', 'teacher_id', 'course_id', 'text', 'default_points'],
                [cls.question_in_id, cls.teacher_in_id, cls.course_in_id, 'سوال داخل محدوده', 2])
        cls.question_in2_id = str(uuid.uuid4())
        _insert('questions', ['id', 'teacher_id', 'course_id', 'text', 'default_points'],
                [cls.question_in2_id, cls.teacher_in_id, cls.course_in_id, 'آ - سوال دوم داخل محدوده', 3])
        cls.question_out_id = str(uuid.uuid4())
        _insert('questions', ['id', 'teacher_id', 'course_id', 'text', 'default_points'],
                [cls.question_out_id, cls.teacher_out_id, cls.course_out_id, 'سوال خارج محدوده', 2])

        cls.admin_id = cls._make_profile('admin_test', 'مدیر سیستم تست')
        _insert('user_roles', ['id', 'user_id', 'role'], [str(uuid.uuid4()), cls.admin_id, 'admin'])

        User = get_user_model()
        cls.manager_user = User.objects.create_user(username='manager_test', email='manager_test@example.com', password='x')
        cls.admin_user = User.objects.create_user(username='admin_test', email='admin_test@example.com', password='x')

    @staticmethod
    def _make_profile(username, full_name):
        profile_id = str(uuid.uuid4())
        _insert('profiles', ['id', 'full_name', 'username', 'email'], [profile_id, full_name, username, f'{username}@example.com'])
        return profile_id


class ExamManagerScopingTests(_ScopedOrgFixture):
    """
    مدیر آموزشی فقط باید اساتید و دانشجویانِ واحد سازمانیِ خودش را ببیند،
    نه همه‌ی کاربران سیستم را. این تست رگرسیون همان باگ دسترسی است که در
    exam_manager_users/courses/groups رفع شد.
    """

    def test_manager_sees_only_own_org_unit_users(self):
        self.client.force_login(self.manager_user)
        response = self.client.get(reverse('core:exam_manager_users'))
        self.assertEqual(response.status_code, 200)
        student_ids = {row['id'] for row in response.context['students']}
        teacher_ids = {row['id'] for row in response.context['teachers']}
        self.assertIn(self.student_in_id, student_ids)
        self.assertNotIn(self.student_out_id, student_ids)
        self.assertIn(self.teacher_in_id, teacher_ids)
        self.assertNotIn(self.teacher_out_id, teacher_ids)

    def test_manager_sees_only_own_org_unit_courses_and_groups(self):
        self.client.force_login(self.manager_user)
        courses_response = self.client.get(reverse('core:exam_manager_courses'))
        self.assertEqual(courses_response.status_code, 200)
        course_ids = {row['id'] for row in courses_response.context['rows']}
        self.assertIn(self.course_in_id, course_ids)
        self.assertNotIn(self.course_out_id, course_ids)

        groups_response = self.client.get(reverse('core:exam_manager_groups'))
        self.assertEqual(groups_response.status_code, 200)
        group_ids = {row['id'] for row in groups_response.context['rows']}
        self.assertIn(self.group_in_id, group_ids)
        self.assertNotIn(self.group_out_id, group_ids)

    def test_admin_still_sees_everyone_system_wide(self):
        self.client.force_login(self.admin_user)
        response = self.client.get(reverse('core:exam_manager_users'))
        self.assertEqual(response.status_code, 200)
        student_ids = {row['id'] for row in response.context['students']}
        teacher_ids = {row['id'] for row in response.context['teachers']}
        self.assertIn(self.student_in_id, student_ids)
        self.assertIn(self.student_out_id, student_ids)
        self.assertIn(self.teacher_in_id, teacher_ids)
        self.assertIn(self.teacher_out_id, teacher_ids)


class SuperAdminUsersUnificationTests(_ScopedOrgFixture):
    """
    صفحه‌ی کاربران مدیر آموزشی و مدیر سیستم اکنون یک صفحه‌ی مشترک است
    (super_admin_users) که فقط سطح دسترسی را بر اساس نقش محدود می‌کند:
    مدیر آموزشی فقط اساتید/دانشجویان واحد خودش را می‌بیند و تب «مدیران»
    و مدیریت سایر مدیران را ندارد؛ عملیات فعال/غیرفعال/حذف/پروفایل هم
    سمت سرور به همین محدوده گیت شده‌اند، نه فقط در UI.
    """

    def test_exam_manager_and_super_admin_urls_show_identical_scoped_content(self):
        self.client.force_login(self.manager_user)
        for url_name in ('core:exam_manager_users', 'core:super_admin_users'):
            response = self.client.get(reverse(url_name))
            self.assertEqual(response.status_code, 200)
            student_ids = {row['id'] for row in response.context['students']}
            teacher_ids = {row['id'] for row in response.context['teachers']}
            self.assertIn(self.student_in_id, student_ids)
            self.assertNotIn(self.student_out_id, student_ids)
            self.assertIn(self.teacher_in_id, teacher_ids)
            self.assertNotIn(self.teacher_out_id, teacher_ids)

    def test_manager_cannot_see_or_manage_other_managers(self):
        self.client.force_login(self.manager_user)
        response = self.client.get(reverse('core:super_admin_users'), {'tab': 'managers'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['active_tab'], 'students')
        self.assertEqual(list(response.context['managers']), [])

        save_response = self.client.post(reverse('core:super_admin_users'), {'admin_action': 'save', 'full_name': 'نفوذی'})
        self.assertEqual(save_response.status_code, 403)

    def test_manager_blocked_from_mutating_out_of_scope_accounts(self):
        self.client.force_login(self.manager_user)
        toggle_url = reverse('core:super_admin_toggle_account_status', args=['student', self.student_out_id])
        self.assertEqual(self.client.post(toggle_url, {'action': 'deactivate'}).status_code, 403)

        delete_url = reverse('core:super_admin_delete_account', args=['teacher', self.teacher_out_id])
        self.assertEqual(self.client.post(delete_url, {}).status_code, 403)

        manager_toggle_url = reverse('core:super_admin_toggle_account_status', args=['manager', self.admin_id])
        self.assertEqual(self.client.post(manager_toggle_url, {'action': 'deactivate'}).status_code, 403)

        profile_url = reverse('core:super_admin_user_profile', args=['student', self.student_out_id])
        self.assertEqual(self.client.get(profile_url).status_code, 404)

        manager_profile_url = reverse('core:super_admin_user_profile', args=['manager', self.admin_id])
        self.assertEqual(self.client.get(manager_profile_url).status_code, 404)

    def test_manager_can_mutate_and_view_in_scope_accounts(self):
        self.client.force_login(self.manager_user)
        toggle_url = reverse('core:super_admin_toggle_account_status', args=['student', self.student_in_id])
        response = self.client.post(toggle_url, {'action': 'deactivate'}, follow=True)
        self.assertEqual(response.status_code, 200)

        profile_url = reverse('core:super_admin_user_profile', args=['student', self.student_in_id])
        self.assertEqual(self.client.get(profile_url).status_code, 200)


class SuperAdminCoursesUnificationTests(_ScopedOrgFixture):
    """
    صفحه‌ی درس‌های مدیر آموزشی و مدیر سیستم هم اکنون یک صفحه‌ی مشترک است
    (super_admin_courses/super_admin_course_form). قبلاً exam_manager_courses
    به‌اشتباه لیست گروه‌ها را نشان می‌داد، نه لیست درس‌ها؛ این هم در همین
    یکسان‌سازی رفع شد.
    """

    def test_exam_manager_and_super_admin_course_urls_show_identical_scoped_content(self):
        self.client.force_login(self.manager_user)
        for url_name in ('core:exam_manager_courses', 'core:super_admin_courses'):
            response = self.client.get(reverse(url_name))
            self.assertEqual(response.status_code, 200)
            course_ids = {row['id'] for row in response.context['rows']}
            self.assertIn(self.course_in_id, course_ids)
            self.assertNotIn(self.course_out_id, course_ids)

    def test_manager_blocked_from_mutating_out_of_scope_course(self):
        self.client.force_login(self.manager_user)
        courses_url = reverse('core:super_admin_courses')

        delete_response = self.client.post(courses_url, {'course_action': 'delete', 'course_id': self.course_out_id})
        self.assertEqual(delete_response.status_code, 403)

        save_response = self.client.post(courses_url, {
            'course_action': 'save',
            'course_id': self.course_out_id,
            'title': 'دستکاری',
            'department_id': self.ou_scoped,
        })
        self.assertEqual(save_response.status_code, 403)

        hijack_response = self.client.post(courses_url, {
            'course_action': 'save',
            'title': 'درس جدید با واحد غیرمجاز',
            'department_id': self.ou_other,
        })
        self.assertEqual(hijack_response.status_code, 403)

    def test_manager_can_manage_in_scope_course(self):
        self.client.force_login(self.manager_user)
        courses_url = reverse('core:super_admin_courses')
        response = self.client.post(courses_url, {
            'course_action': 'save',
            'title': 'درس جدید داخل محدوده',
            'department_id': self.ou_scoped,
        })
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['ok'])


class SuperAdminGroupsUnificationTests(_ScopedOrgFixture):
    """
    صفحه‌ی گروه‌بندی مدیر آموزشی و مدیر سیستم هم اکنون یک صفحه‌ی مشترک
    است (super_admin_groups/super_admin_group_create/super_admin_group_edit).
    """

    def test_exam_manager_and_super_admin_group_urls_show_identical_scoped_content(self):
        self.client.force_login(self.manager_user)
        for url_name in ('core:exam_manager_groups', 'core:super_admin_groups'):
            response = self.client.get(reverse(url_name))
            self.assertEqual(response.status_code, 200)
            group_ids = {row['id'] for row in response.context['rows']}
            self.assertIn(self.group_in_id, group_ids)
            self.assertNotIn(self.group_out_id, group_ids)

    def test_manager_cannot_view_or_edit_out_of_scope_group(self):
        self.client.force_login(self.manager_user)
        edit_url = reverse('core:super_admin_group_edit', args=[self.group_out_id])
        self.assertEqual(self.client.get(edit_url).status_code, 404)

        detail_url = reverse('core:exam_manager_group_detail', args=[self.group_out_id])
        self.assertEqual(self.client.get(detail_url).status_code, 404)

    def test_manager_blocked_from_assigning_out_of_scope_course_or_teacher(self):
        self.client.force_login(self.manager_user)
        groups_url = reverse('core:super_admin_groups')

        delete_response = self.client.post(groups_url, {'group_action': 'delete', 'group_id': self.group_out_id})
        self.assertEqual(delete_response.status_code, 403)

        hijack_response = self.client.post(groups_url, {
            'group_action': 'save',
            'course_id': self.course_out_id,
            'teacher_id': self.teacher_out_id,
            'course_name': 'گروه نفوذی',
        })
        self.assertEqual(hijack_response.status_code, 403)

        edit_response = self.client.post(reverse('core:super_admin_group_edit', args=[self.group_in_id]), {
            'course_id': self.course_out_id,
            'teacher_ids': [self.teacher_out_id],
        })
        self.assertEqual(edit_response.status_code, 403)

    def test_manager_can_manage_in_scope_group(self):
        self.client.force_login(self.manager_user)
        groups_url = reverse('core:super_admin_groups')
        response = self.client.post(groups_url, {
            'group_action': 'save',
            'course_id': self.course_in_id,
            'teacher_id': self.teacher_in_id,
            'course_name': 'گروه جدید داخل محدوده',
        }, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['ok'])


class SuperAdminExamsUnificationTests(_ScopedOrgFixture):
    """
    صفحه‌ی آزمون‌های مدیر آموزشی و مدیر سیستم هم اکنون یک صفحه‌ی مشترک
    است (super_admin_exams/super_admin_exam_detail/super_admin_exam_edit).
    این تست همچنین حفره‌ی scope مستقلی را پوشش می‌دهد که در _em_exam_rows
    پیدا شد: قبلاً هر مدیر آموزشی همه‌ی آزمون‌های کل سیستم را می‌دید.
    """

    def test_exam_manager_and_super_admin_exam_urls_show_identical_scoped_content(self):
        self.client.force_login(self.manager_user)
        for url_name in ('core:exam_manager_exams', 'core:super_admin_exams'):
            response = self.client.get(reverse(url_name))
            self.assertEqual(response.status_code, 200)
            exam_ids = {row['id'] for row in response.context['rows']}
            self.assertIn(self.exam_in_id, exam_ids)
            self.assertNotIn(self.exam_out_id, exam_ids)

    def test_admin_still_sees_every_exam(self):
        self.client.force_login(self.admin_user)
        response = self.client.get(reverse('core:super_admin_exams'))
        exam_ids = {row['id'] for row in response.context['rows']}
        self.assertIn(self.exam_in_id, exam_ids)
        self.assertIn(self.exam_out_id, exam_ids)

    def test_manager_cannot_view_out_of_scope_exam(self):
        self.client.force_login(self.manager_user)
        self.assertEqual(self.client.get(reverse('core:super_admin_exam_detail', args=[self.exam_out_id])).status_code, 404)
        self.assertEqual(self.client.get(reverse('core:super_admin_exam_edit', args=[self.exam_out_id])).status_code, 404)
        self.assertEqual(self.client.get(reverse('core:exam_manager_exam_questions', args=[self.exam_out_id])).status_code, 404)

    def test_manager_can_view_in_scope_exam(self):
        self.client.force_login(self.manager_user)
        self.assertEqual(self.client.get(reverse('core:super_admin_exam_detail', args=[self.exam_in_id])).status_code, 200)
        self.assertEqual(self.client.get(reverse('core:super_admin_exam_edit', args=[self.exam_in_id])).status_code, 200)

    def test_manager_blocked_from_creating_exam_with_out_of_scope_group(self):
        self.client.force_login(self.manager_user)
        response = self.client.post(reverse('core:super_admin_exams'), {
            'exam_action': 'create',
            'group_id': self.group_out_id,
            'title': 'آزمون نفوذی',
        })
        self.assertEqual(response.status_code, 403)


class SuperAdminCalendarUnificationTests(_ScopedOrgFixture):
    """
    تقویم مدیر آموزشی و مدیر سیستم هم اکنون یک صفحه‌ی مشترک است
    (super_admin_calendar) که از تبدیل واقعی میلادی-به-جلالی استفاده
    می‌کند، نه محاسبه‌ی جعلی/هاردکد قبلی exam_manager (`is_today = day
    == 25` و ...). این هم اکنون حفره‌ی scope مستقل _em_calendar_rows را
    پوشش می‌دهد: قبلاً تقویم مدیر آموزشی همه‌ی آزمون‌ها و رویدادهای
    تقویمی کل سیستم را نشان می‌داد.
    """

    def test_exam_manager_and_super_admin_calendar_urls_show_identical_scoped_content(self):
        self.client.force_login(self.manager_user)
        for url_name in ('core:exam_manager_calendar', 'core:super_admin_calendar'):
            response = self.client.get(reverse(url_name))
            self.assertEqual(response.status_code, 200)
            event_ids = {event['id'] for event in response.context['events']}
            self.assertIn(self.exam_in_id, event_ids)
            self.assertNotIn(self.exam_out_id, event_ids)

    def test_admin_calendar_shows_every_exam(self):
        self.client.force_login(self.admin_user)
        response = self.client.get(reverse('core:super_admin_calendar'))
        event_ids = {event['id'] for event in response.context['events']}
        self.assertIn(self.exam_in_id, event_ids)
        self.assertIn(self.exam_out_id, event_ids)


class SuperAdminExamBulkImportScopeTests(_ScopedOrgFixture):
    """
    ویزارد ورود گروهی آزمون قبلاً کد درس/استاد/گروه را در کل سیستم
    تطبیق می‌داد (بدون scope) و برای هیچ نقشی جز admin در دسترس نبود.
    این تست تایید می‌کند مدیر آموزشی هم اکنون به آن دسترسی دارد، اما
    فقط می‌تواند به درس/استاد داخل محدوده‌ی خودش وصل شود.
    """

    def _set_draft(self, rows):
        session = self.client.session
        session['super_admin_exam_bulk_import'] = {
            'step': 3,
            'headers': list(rows[0].keys()),
            'rows': rows,
            'mapping': {
                'title': 'title', 'course_code': 'course_code', 'teacher_code': 'teacher_code',
                'group_code': 'group_code', 'date': 'date', 'start_time': 'start_time',
                'duration': 'duration', 'passing_score': 'passing_score',
            },
        }
        session.save()

    def test_manager_can_reach_wizard(self):
        self.client.force_login(self.manager_user)
        self.assertEqual(self.client.get(reverse('core:super_admin_exam_bulk_import')).status_code, 200)

    def test_manager_row_matching_in_scope_course_and_teacher_resolves_ok(self):
        self.client.force_login(self.manager_user)
        self._set_draft([{
            'title': 'آزمون وارداتی', 'course_code': 'درس داخل محدوده', 'teacher_code': 'استاد داخل محدوده',
            'group_code': '01', 'date': '1404/03/25', 'start_time': '09:00', 'duration': '90', 'passing_score': '10',
        }])
        response = self.client.get(reverse('core:super_admin_exam_bulk_import'), {'step': 3})
        record = response.context['records'][0]
        self.assertIn(record['level'], {'ok', 'warning'})
        self.assertEqual(record['course_id'], self.course_in_id)
        self.assertEqual(record['teacher_id'], self.teacher_in_id)

    def test_manager_row_matching_out_of_scope_course_and_teacher_fails(self):
        self.client.force_login(self.manager_user)
        self._set_draft([{
            'title': 'آزمون نفوذی', 'course_code': 'درس خارج محدوده', 'teacher_code': 'استاد خارج محدوده',
            'group_code': '01', 'date': '1404/03/25', 'start_time': '09:00', 'duration': '90', 'passing_score': '10',
        }])
        response = self.client.get(reverse('core:super_admin_exam_bulk_import'), {'step': 3})
        record = response.context['records'][0]
        self.assertEqual(record['level'], 'error')
        self.assertIn('کد درس در سامانه پیدا نشد.', record['issues'])
        self.assertIn('کد استاد در سامانه پیدا نشد.', record['issues'])
        self.assertEqual(record['course_id'], '')
        self.assertEqual(record['teacher_id'], '')

    def test_admin_row_matching_out_of_scope_course_and_teacher_resolves_ok(self):
        self.client.force_login(self.admin_user)
        self._set_draft([{
            'title': 'آزمون سراسری', 'course_code': 'درس خارج محدوده', 'teacher_code': 'استاد خارج محدوده',
            'group_code': '01', 'date': '1404/03/25', 'start_time': '09:00', 'duration': '90', 'passing_score': '10',
        }])
        response = self.client.get(reverse('core:super_admin_exam_bulk_import'), {'step': 3})
        record = response.context['records'][0]
        self.assertIn(record['level'], {'ok', 'warning'})
        self.assertEqual(record['course_id'], self.course_out_id)
        self.assertEqual(record['teacher_id'], self.teacher_out_id)


class SuperAdminExamWizardTests(_ScopedOrgFixture):
    """
    ویزارد ۷ مرحله‌ای ایجاد آزمون: منابع (درس/استاد/دانشجو/گروه/سوال) باید
    برای مدیر آموزشی scope شوند و ثبت نهایی هم سمت سرور دوباره بررسی شود،
    نه فقط مخفی‌کردن گزینه‌های خارج از محدوده در فرم.
    """

    def _base_payload(self, **overrides):
        payload = {
            'title': 'آزمون تست',
            'course_id': self.course_in_id,
            'teacher_id': self.teacher_in_id,
            'group_ids': [self.group_in_id],
            'question_ids': [self.question_in_id],
            'duration_minutes': '90',
            'start_at': '2026-09-10T09:00',
            'publish_choice': 'draft',
        }
        payload.update(overrides)
        return payload

    def test_manager_sees_only_in_scope_resources_on_get(self):
        self.client.force_login(self.manager_user)
        response = self.client.get(reverse('core:super_admin_exam_create'))
        self.assertEqual(response.status_code, 200)
        course_ids = {c['id'] for c in response.context['courses']}
        teacher_ids = {t['id'] for t in response.context['teachers']}
        question_ids = {q['id'] for q in response.context['questions']}
        group_ids = {g['id'] for g in response.context['groups']}
        self.assertIn(self.course_in_id, course_ids)
        self.assertNotIn(self.course_out_id, course_ids)
        self.assertIn(self.teacher_in_id, teacher_ids)
        self.assertNotIn(self.teacher_out_id, teacher_ids)
        self.assertIn(self.question_in_id, question_ids)
        self.assertNotIn(self.question_out_id, question_ids)
        self.assertIn(self.group_in_id, group_ids)
        self.assertNotIn(self.group_out_id, group_ids)

    def test_manager_can_create_exam_with_in_scope_resources_and_staff(self):
        self.client.force_login(self.manager_user)
        payload = self._base_payload(staff_teacher_id=[self.teacher_in_id], staff_role=['designer'])
        response = self.client.post(reverse('core:super_admin_exam_create'), payload, follow=True)
        self.assertEqual(response.status_code, 200)
        with connection.cursor() as cursor:
            cursor.execute("SELECT id, security_level, execution_mode, exam_mode FROM exams WHERE title = 'آزمون تست'")
            row = cursor.fetchone()
        self.assertIsNotNone(row)
        exam_id, security_level, execution_mode, exam_mode = row
        self.assertEqual(security_level, 'medium')
        self.assertEqual(execution_mode, 'web')
        self.assertEqual(exam_mode, 'standard')
        with connection.cursor() as cursor:
            cursor.execute('SELECT teacher_id, role FROM exam_staff WHERE exam_id = %s', [exam_id])
            staff_rows = cursor.fetchall()
        self.assertEqual(staff_rows, [(self.teacher_in_id, 'designer')])

    def test_exam_questions_preserve_submitted_order_not_alphabetical(self):
        """
        ترتیب چینش سوالات (drag-reorder در ویزارد) باید همان ترتیبی باشد که کاربر
        submit کرده، نه ترتیب الفبایی متن سوال - چون قبلاً کوئری با ORDER BY text
        این ترتیب را نادیده می‌گرفت.
        """
        self.client.force_login(self.manager_user)
        payload = self._base_payload(question_ids=[self.question_in_id, self.question_in2_id])
        response = self.client.post(reverse('core:super_admin_exam_create'), payload, follow=True)
        self.assertEqual(response.status_code, 200)
        with connection.cursor() as cursor:
            cursor.execute("SELECT id FROM exams WHERE title = 'آزمون تست'")
            exam_id = cursor.fetchone()[0]
            cursor.execute('SELECT question_id FROM exam_questions WHERE exam_id = %s ORDER BY order_index', [exam_id])
            ordered_question_ids = [row[0] for row in cursor.fetchall()]
        self.assertEqual(ordered_question_ids, [self.question_in_id, self.question_in2_id])

    def test_manager_blocked_from_using_out_of_scope_course(self):
        self.client.force_login(self.manager_user)
        payload = self._base_payload(course_id=self.course_out_id)
        response = self.client.post(reverse('core:super_admin_exam_create'), payload)
        self.assertEqual(response.status_code, 403)

    def test_manager_blocked_from_using_out_of_scope_question(self):
        self.client.force_login(self.manager_user)
        payload = self._base_payload(question_ids=[self.question_out_id])
        response = self.client.post(reverse('core:super_admin_exam_create'), payload)
        self.assertEqual(response.status_code, 403)

    def test_manager_blocked_from_assigning_out_of_scope_staff(self):
        self.client.force_login(self.manager_user)
        payload = self._base_payload(staff_teacher_id=[self.teacher_out_id], staff_role=['observer'])
        response = self.client.post(reverse('core:super_admin_exam_create'), payload)
        self.assertEqual(response.status_code, 403)

    def test_manager_blocked_from_assigning_out_of_scope_group_or_student(self):
        self.client.force_login(self.manager_user)
        for payload in (
            self._base_payload(group_ids=[self.group_in_id, self.group_out_id]),
            self._base_payload(group_ids=[], student_ids=[self.student_out_id]),
            self._base_payload(excluded_student_ids=[self.student_out_id]),
        ):
            response = self.client.post(reverse('core:super_admin_exam_create'), payload)
            self.assertEqual(response.status_code, 403)
        with connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) FROM exams WHERE title = 'آزمون تست'")
            self.assertEqual(cursor.fetchone()[0], 0)

    def test_manager_can_assign_in_scope_student_directly(self):
        self.client.force_login(self.manager_user)
        payload = self._base_payload(group_ids=[], student_ids=[self.student_in_id])
        response = self.client.post(reverse('core:super_admin_exam_create'), payload)
        self.assertEqual(response.status_code, 302)

    def test_admin_sees_all_resources_and_can_use_any(self):
        self.client.force_login(self.admin_user)
        response = self.client.get(reverse('core:super_admin_exam_create'))
        course_ids = {c['id'] for c in response.context['courses']}
        self.assertIn(self.course_in_id, course_ids)
        self.assertIn(self.course_out_id, course_ids)
        payload = self._base_payload(title='آزمون ادمین', course_id=self.course_out_id, teacher_id=self.teacher_out_id, question_ids=[self.question_out_id])
        response = self.client.post(reverse('core:super_admin_exam_create'), payload, follow=True)
        self.assertEqual(response.status_code, 200)
        with connection.cursor() as cursor:
            cursor.execute("SELECT 1 FROM exams WHERE title = 'آزمون ادمین'")
            self.assertIsNotNone(cursor.fetchone())


class SuperAdminExamParticipantsMatchTests(_ScopedOrgFixture):
    """تطبیق فایل مخاطبان آزمون هم باید مثل تطبیق دانشجویان bulk-import محدود به scope مدیر باشد."""

    NID_IN = 'NID-IN-001'
    NID_OUT = 'NID-OUT-002'

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        with connection.cursor() as cursor:
            cursor.execute('UPDATE profiles SET national_id = %s WHERE id = %s', [cls.NID_IN, cls.student_in_id])
            cursor.execute('UPDATE profiles SET national_id = %s WHERE id = %s', [cls.NID_OUT, cls.student_out_id])

    def _upload_csv(self, url, identifier):
        content = f'کد ملی یا شماره دانشجویی\n{identifier}\n'.encode('utf-8-sig')
        upload = io.BytesIO(content)
        upload.name = 'participants.csv'
        return self.client.post(url, {'excel_file': upload})

    def test_manager_can_match_in_scope_student(self):
        self.client.force_login(self.manager_user)
        response = self._upload_csv(reverse('core:super_admin_exam_participants_match'), self.NID_IN)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(len(data['found']), 1)
        self.assertEqual(data['found'][0]['id'], self.student_in_id)

    def test_manager_cannot_match_out_of_scope_student(self):
        self.client.force_login(self.manager_user)
        response = self._upload_csv(reverse('core:super_admin_exam_participants_match'), self.NID_OUT)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(len(data['found']), 0)
        self.assertEqual(data['not_found'], [self.NID_OUT])

    def test_admin_can_match_any_student(self):
        self.client.force_login(self.admin_user)
        response = self._upload_csv(reverse('core:super_admin_exam_participants_match'), self.NID_OUT)
        data = response.json()
        self.assertEqual(len(data['found']), 1)
        self.assertEqual(data['found'][0]['id'], self.student_out_id)


class TeacherQuestionBuilderTests(_ScopedOrgFixture):
    """
    فرم ایجاد/ویرایش سؤال استاد: قبلاً داده‌ی پاسخ از فیلدهای مخفیِ هاردکد
    (seed) ارسال می‌شد، course_id اصلاً ذخیره نمی‌شد، و صفحه‌ی ویرایش یک
    ماکاپ بی‌عملکرد بود. این تست‌ها ذخیره‌ی واقعی، اعتبارسنجی سمت سرور،
    محدودیت درس و مالکیت سؤال در ویرایش را پوشش می‌دهند.
    """

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        User = get_user_model()
        cls.teacher_user = User.objects.create_user(username='teacher_in_scope', email='t_in@example.com', password='x')
        cls.other_teacher_user = User.objects.create_user(username='teacher_out_scope', email='t_out@example.com', password='x')

    def _create(self, qtype, data):
        url = f"{reverse('core:teacher_question_create')}?type={qtype}"
        payload = {'type': qtype, 'structure': 'independent', 'course_id': self.course_in_id, 'is_published': '1'}
        payload.update(data)
        return self.client.post(url, payload)

    def _row(self, text):
        with connection.cursor() as cursor:
            cursor.execute(
                'SELECT id, type, course_id, options, correct_answer, matching_pairs, ordering_items, teacher_id FROM questions WHERE text = %s',
                [text],
            )
            return cursor.fetchone()

    def test_matching_question_saves_real_pairs_and_course(self):
        self.client.force_login(self.teacher_user)
        response = self._create('matching', {
            'text': 'سوال تطبیقی تست',
            'matching_left': ['الف۱', 'الف۲'],
            'matching_right': ['ب۱', 'ب۲'],
        })
        self.assertEqual(response.status_code, 302)
        row = self._row('سوال تطبیقی تست')
        self.assertIsNotNone(row)
        self.assertEqual(row[1], 'matching')
        self.assertEqual(row[2], self.course_in_id)
        pairs = json.loads(row[5])
        self.assertEqual([(p['left'], p['right']) for p in pairs], [('الف۱', 'ب۱'), ('الف۲', 'ب۲')])

    def test_fill_blank_saves_typed_alternatives_per_blank(self):
        self.client.force_login(self.teacher_user)
        text = 'جمجمه از [۱] و [۲] تشکیل شده است.'
        response = self._create('fill_blank', {
            'text': text,
            'blank_count': '2',
            'blank_answers_1': 'نوروکرانیوم، neurocranium',
            'blank_points_1': '2',
            'blank_answers_2': 'ویسروکرانیوم',
            'blank_spaces_2': '1',
        })
        self.assertEqual(response.status_code, 302)
        row = self._row(text)
        self.assertEqual(json.loads(row[4]), 'نوروکرانیوم|ویسروکرانیوم')
        self.assertEqual(json.loads(row[3]), [])
        with connection.cursor() as cursor:
            cursor.execute('SELECT accepted_answers FROM questions WHERE id = %s', [row[0]])
            blanks = json.loads(cursor.fetchone()[0])
        self.assertEqual(blanks[0]['answers'], ['نوروکرانیوم', 'neurocranium'])
        self.assertEqual(blanks[0]['points'], 2)
        self.assertTrue(blanks[1]['ignore_spaces'])

    def test_scoring_settings_are_saved(self):
        self.client.force_login(self.teacher_user)
        response = self._create('multi', {
            'text': 'چندپاسخی با تنظیمات',
            'option_text': ['الف', 'ب', 'ج', 'د'],
            'correct_options': ['1', '2'],
            'scoring_mode': 'partial',
            'min_correct': '1',
            'max_correct': '3',
            'floor_zero': '1',
            'negative_marking': '1',
            'negative_value': '۰.۲۵',
            'shuffle_options': '1',
        })
        self.assertEqual(response.status_code, 302)
        with connection.cursor() as cursor:
            cursor.execute('SELECT scoring_settings, negative_points FROM questions WHERE text = %s', ['چندپاسخی با تنظیمات'])
            settings_json, negative_points = cursor.fetchone()
        settings_data = json.loads(settings_json)
        self.assertEqual(settings_data['scoring_mode'], 'partial')
        self.assertEqual(settings_data['max_correct'], 3)
        self.assertTrue(settings_data['shuffle_options'])
        self.assertEqual(settings_data['negative_value'], 0.25)
        self.assertEqual(float(negative_points), -0.25)

    def test_multi_max_selection_must_cover_correct_options(self):
        self.client.force_login(self.teacher_user)
        response = self._create('multi', {
            'text': 'چندپاسخی ناسازگار',
            'option_text': ['الف', 'ب', 'ج'],
            'correct_options': ['1', '2', '3'],
            'max_correct': '2',
        })
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(self._row('چندپاسخی ناسازگار'))

    def test_multi_form_layout_and_all_or_nothing_choice(self):
        self.client.force_login(self.teacher_user)
        page = self.client.get(f"{reverse('core:teacher_question_create')}?type=multi")
        self.assertContains(page, 'class="qb qb--mc qb--v2"')
        self.assertContains(page, 'name="scoring_mode" value="all_or_nothing"')
        self.assertNotContains(page, '<select data-template-select>')  # template pills replace the side select
        self.assertContains(page, 'data-penalty-all')
        response = self._create('multi', {
            'text': 'چندپاسخی همه یا هیچ', 'option_text': ['الف', 'ب', 'ج'], 'correct_options': ['1', '2'],
            'scoring_mode': 'all_or_nothing', 'min_correct': '1', 'max_correct': '3',
        })
        self.assertEqual(response.status_code, 302)
        with connection.cursor() as cursor:
            cursor.execute("SELECT scoring_settings FROM questions WHERE text = %s", ['چندپاسخی همه یا هیچ'])
            self.assertEqual(json.loads(cursor.fetchone()[0])['scoring_mode'], 'all_or_nothing')

    def test_short_answer_invalid_regex_is_rejected(self):
        self.client.force_login(self.teacher_user)
        response = self._create('short_answer', {'text': 'regex نامعتبر', 'accepted_answer': ['(abc'], 'accepted_mode_1': 'regex'})
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(self._row('regex نامعتبر'))

    def test_short_answer_saves_typo_word_order_and_scoring_settings(self):
        self.client.force_login(self.teacher_user)
        response = self._create('short_answer', {
            'text': 'پاسخ کوتاه با قواعد جدید',
            'accepted_answer': ['گاستریت هلیکوباکتر', 'H. pylori gastritis'],
            'accepted_points_1': '100', 'accepted_mode_1': 'exact', 'accepted_typo_1': '1', 'accepted_word_order_1': '1',
            'accepted_points_2': '50', 'accepted_mode_2': 'contains', 'accepted_typo_2': '9',
            'scoring_mode': ['all_or_nothing', 'partial'],
            'correct_threshold': '50',
        })
        self.assertEqual(response.status_code, 302)
        with connection.cursor() as cursor:
            cursor.execute('SELECT accepted_answers, scoring_settings FROM questions WHERE text = %s', ['پاسخ کوتاه با قواعد جدید'])
            accepted, settings_data = (json.loads(v) if isinstance(v, str) else v for v in cursor.fetchone())
        self.assertEqual(accepted[0]['typo_tolerance'], 1)
        self.assertTrue(accepted[0]['ignore_word_order'])
        self.assertEqual(accepted[1]['typo_tolerance'], 2)  # clamped
        self.assertFalse(accepted[1]['ignore_word_order'])
        self.assertEqual(settings_data['scoring_mode'], 'partial')
        self.assertEqual(settings_data['correct_threshold'], 50)

    def test_question_image_upload_is_validated_and_stored(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.force_login(self.teacher_user)
        with self.settings(MEDIA_ROOT=tempfile.mkdtemp()):
            bad = SimpleUploadedFile('payload.exe', b'MZ', content_type='application/octet-stream')
            response = self._create('single', {'text': 'سوال با فایل بد', 'option_text': ['الف', 'ب'], 'correct_options': ['1'], 'question_image_file': bad})
            self.assertEqual(response.status_code, 200)
            self.assertIsNone(self._row('سوال با فایل بد'))
            image = SimpleUploadedFile('ecg.png', b'\x89PNG\r\n\x1a\n' + b'0' * 32, content_type='image/png')
            response = self._create('single', {
                'text': 'سوال با تصویر', 'option_text': ['الف', 'ب'], 'correct_options': ['1'],
                'question_image_file': image, 'question_image_alt': 'نوار قلب',
            })
            self.assertEqual(response.status_code, 302)
        with connection.cursor() as cursor:
            cursor.execute('SELECT question_media FROM questions WHERE text = %s', ['سوال با تصویر'])
            media = json.loads(cursor.fetchone()[0])
        self.assertTrue(media['image'].startswith('/media/question-media/') or 'question-media/' in media['image'])
        self.assertTrue(media['image'].endswith('.png'))
        self.assertEqual(media['image_alt'], 'نوار قلب')

    def test_fill_blank_media_hint_is_stored_reloaded_and_validated(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from apps.core.views import _question_form_data
        self.client.force_login(self.teacher_user)
        base = {'text': 'جمجمه [۱] و [۲]', 'blank_count': '2', 'blank_answers_1': 'الف', 'blank_answers_2': 'ب'}
        with self.settings(MEDIA_ROOT=tempfile.mkdtemp()):
            bad = SimpleUploadedFile('hint.exe', b'MZ', content_type='application/octet-stream')
            self.assertEqual(self._create('fill_blank', dict(base, blank_media_file_2=bad)).status_code, 200)
            self.assertIsNone(self._row('جمجمه [۱] و [۲]'))
            image = SimpleUploadedFile('skull.png', b'\x89PNG\r\n\x1a\n' + b'0' * 32, content_type='image/png')
            self.assertEqual(self._create('fill_blank', dict(base, blank_media_file_2=image)).status_code, 302)
        with connection.cursor() as cursor:
            cursor.execute('SELECT accepted_answers, type, text FROM questions WHERE text = %s', ['جمجمه [۱] و [۲]'])
            accepted, kind, text = cursor.fetchone()
        blanks = json.loads(accepted)
        self.assertEqual(blanks[0]['media'], '')
        self.assertIn('question-media/', blanks[1]['media'])
        self.assertEqual(blanks[1]['media_kind'], 'image')
        form = _question_form_data({'type': kind, 'text': text, 'accepted_answers': accepted})
        self.assertEqual([blank['media_kind'] for blank in form['blanks']], ['', 'image'])

    def test_video_subtitle_is_validated_and_needs_a_video(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.force_login(self.teacher_user)
        base = {'option_text': ['الف', 'ب'], 'correct_options': ['1']}
        video = lambda: SimpleUploadedFile('case.mp4', b'\x00\x00\x00\x18ftypmp42' + b'0' * 32, content_type='video/mp4')
        vtt = lambda body=b'\xef\xbb\xbfWEBVTT\n\n00:00.000 --> 00:02.000\n\xd8\xb3\xd9\x84\xd8\xa7\xd9\x85\n': SimpleUploadedFile('sub.vtt', body, content_type='text/vtt')
        with self.settings(MEDIA_ROOT=tempfile.mkdtemp()):
            fake = vtt(b'<script>alert(1)</script>')
            response = self._create('single', dict(base, text='زیرنویس جعلی', question_video_file=video(), question_video_subtitle_file=fake))
            self.assertEqual(response.status_code, 200)
            self.assertIsNone(self._row('زیرنویس جعلی'))
            self.assertEqual(self._create('single', dict(base, text='با زیرنویس', question_video_file=video(), question_video_subtitle_file=vtt())).status_code, 302)
            self.assertEqual(self._create('single', dict(base, text='زیرنویس بی‌ویدئو', question_video_subtitle_file=vtt())).status_code, 302)
        with connection.cursor() as cursor:
            cursor.execute('SELECT text, question_media FROM questions WHERE text IN (%s, %s)', ['با زیرنویس', 'زیرنویس بی‌ویدئو'])
            media = {text: json.loads(value) for text, value in cursor.fetchall()}
        self.assertTrue(media['با زیرنویس']['video_subtitle'].endswith('.vtt'))
        self.assertEqual(media['زیرنویس بی‌ویدئو']['video_subtitle'], '')

    def test_autosave_stores_private_draft_and_restores_it(self):
        self.client.force_login(self.teacher_user)
        autosave_url = reverse('core:teacher_question_autosave')
        response = self.client.post(autosave_url, {'type': 'ordering', 'text': 'پیش‌نویس خودکار', 'ordering_item': ['الف', 'ب']})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['ok'])
        self.assertIsNone(self._row('پیش‌نویس خودکار'))  # a draft never becomes a question row
        create_url = f"{reverse('core:teacher_question_create')}?type=ordering"
        self.assertTrue(self.client.get(create_url).context['question_draft_saved_at'])
        restored = self.client.get(f'{create_url}&draft=restore').context['qf']
        self.assertEqual(restored['text'], 'پیش‌نویس خودکار')
        self.assertEqual([item['text'] for item in restored['ordering'][:2]], ['الف', 'ب'])
        self._create('ordering', {'text': 'پیش‌نویس خودکار', 'ordering_item': ['الف', 'ب']})
        self.assertNotIn('question_draft_saved_at', self.client.get(create_url).context)

    def test_autosave_rejects_other_teachers_question_and_invalid_type(self):
        self.client.force_login(self.other_teacher_user)
        autosave_url = reverse('core:teacher_question_autosave')
        self.assertEqual(self.client.post(autosave_url, {'type': 'single', 'question_id': self.question_in_id, 'text': 'x'}).status_code, 404)
        self.assertEqual(self.client.post(autosave_url, {'type': 'osce', 'text': 'x'}).status_code, 400)
        self.assertEqual(self.client.get(autosave_url).status_code, 405)

    def test_kept_media_url_outside_question_media_is_dropped(self):
        self.client.force_login(self.teacher_user)
        self._create('single', {
            'text': 'سوال با آدرس دستکاری‌شده', 'option_text': ['الف', 'ب'], 'correct_options': ['1'],
            'existing_question_image': 'https://evil.example/x.png',
        })
        with connection.cursor() as cursor:
            cursor.execute('SELECT question_media FROM questions WHERE text = %s', ['سوال با آدرس دستکاری‌شده'])
            media = json.loads(cursor.fetchone()[0])
        self.assertEqual(media['image'], '')

    def test_no_hardcoded_sample_data_is_saved_for_independent_questions(self):
        self.client.force_login(self.teacher_user)
        self._create('ordering', {'text': 'سوال مرتب‌سازی تست', 'ordering_item': ['اول', 'دوم']})
        with connection.cursor() as cursor:
            cursor.execute('SELECT options, tags, question_media FROM questions WHERE text = %s', ['سوال مرتب‌سازی تست'])
            options, tags, media = cursor.fetchone()
        self.assertEqual(json.loads(options), [])
        self.assertEqual(tags, '')
        self.assertEqual(json.loads(media)['image'], '')

    def test_invalid_matching_is_rejected_and_input_is_kept(self):
        self.client.force_login(self.teacher_user)
        response = self._create('matching', {'text': 'تطبیقی ناقص', 'matching_left': ['فقط چپ'], 'matching_right': ['']})
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(self._row('تطبیقی ناقص'))
        self.assertEqual(response.context['qf']['text'], 'تطبیقی ناقص')
        self.assertEqual(response.context['qf']['pairs'][0]['left'], 'فقط چپ')

    def test_fill_blank_answer_count_must_match_markers(self):
        self.client.force_login(self.teacher_user)
        text = 'دو جای خالی [۱] و [۲]'
        response = self._create('fill_blank', {'text': text, 'blank_count': '1', 'blank_answers_1': 'تنها'})
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(self._row(text))

    def test_out_of_scope_course_is_rejected(self):
        self.client.force_login(self.teacher_user)
        response = self._create('ordering', {'text': 'درس نفوذی', 'ordering_item': ['اول', 'دوم'], 'course_id': self.course_out_id})
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(self._row('درس نفوذی'))

    def test_teacher_can_edit_own_question(self):
        self.client.force_login(self.teacher_user)
        self._create('ordering', {'text': 'قبل از ویرایش', 'ordering_item': ['اول', 'دوم']})
        question_id = self._row('قبل از ویرایش')[0]
        edit_url = reverse('core:teacher_question_edit', args=[question_id])
        get_response = self.client.get(edit_url)
        self.assertEqual(get_response.status_code, 200)
        self.assertEqual([item['text'] for item in get_response.context['qf']['ordering'][:2]], ['اول', 'دوم'])
        response = self.client.post(edit_url, {
            'text': 'بعد از ویرایش',
            'course_id': self.course_in_id,
            'ordering_item': ['دوم', 'اول', 'سوم'],
            'is_published': '1',
        })
        self.assertEqual(response.status_code, 302)
        row = self._row('بعد از ویرایش')
        self.assertEqual(row[0], question_id)
        self.assertEqual([item['text'] for item in json.loads(row[6])], ['دوم', 'اول', 'سوم'])

    def test_teacher_cannot_view_or_edit_another_teachers_question(self):
        self.client.force_login(self.other_teacher_user)
        edit_url = reverse('core:teacher_question_edit', args=[self.question_in_id])
        self.assertEqual(self.client.get(edit_url).status_code, 404)
        self.assertEqual(self.client.post(edit_url, {'text': 'دستکاری', 'type': 'single'}).status_code, 404)
        with connection.cursor() as cursor:
            cursor.execute('SELECT text FROM questions WHERE id = %s', [self.question_in_id])
            self.assertEqual(cursor.fetchone()[0], 'سوال داخل محدوده')

    def test_exam_style_preview_for_every_independent_type(self):
        from apps.core.views import _sx_type_label
        self.client.force_login(self.teacher_user)
        url = reverse('core:teacher_question_create')
        for kind in ('single', 'multi', 'true_false', 'short_answer', 'essay', 'fill_blank', 'matching', 'ordering'):
            html = self.client.get(f'{url}?type={kind}').content.decode()
            self.assertIn('data-exam-preview', html, kind)
            self.assertIn('data-qx-area', html, kind)
            # Same type label as the student exam page.
            self.assertIn(f'سؤال ۱ — {_sx_type_label(kind)}', html, kind)
            self.assertIn('css/exam-view', html, kind)
            self.assertIn('js/exam-interactions', html, kind)
            # The old modal preview is gone.
            self.assertNotIn('data-preview-modal', html, kind)
        multi = self.client.get(f'{url}?type=multi').content.decode()
        self.assertIn('data-qx-guide-variant="max"', multi)
        self.assertIn('حداکثر __N__ گزینه', multi)
        essay = self.client.get(f'{url}?type=essay').content.decode()
        self.assertIn('data-qx-guide-variant="upload"', essay)

    def test_true_false_layout_moves_stem_media_to_side_card(self):
        self.client.force_login(self.teacher_user)
        url = reverse('core:teacher_question_create')
        html = self.client.get(f'{url}?type=true_false').content.decode()
        self.assertIn('qb--tf', html)
        self.assertIn('رسانه سؤال', html)
        self.assertIn('class="qf-drop qb-tf-drop"', html)
        # Stem media fields still render exactly once, so uploads keep their names.
        self.assertEqual(html.count('name="question_video_file"'), 1)
        self.assertIn('<span>صوت</span>', html)
        # Exam-style student preview is included (its markup is covered by the preview tests).
        self.assertIn('data-exam-preview', html)
        single = self.client.get(f'{url}?type=single').content.decode()
        self.assertNotIn('class="qf-drop qb-tf-drop"', single)
        self.assertNotIn('<span>صوت</span>', single)
        self.assertEqual(single.count('name="question_video_file"'), 1)

    def test_fill_blank_layout_shows_media_tiles_and_marker_view(self):
        self.client.force_login(self.teacher_user)
        url = reverse('core:teacher_question_create')
        html = self.client.get(f'{url}?type=fill_blank').content.decode()
        self.assertIn('qb--fb qb--v2', html)
        self.assertIn('class="qf-drop qb-fb-add"', html)
        self.assertIn('تلفظ واژه‌ها (اختیاری)', html)
        self.assertIn('class="qb-fb-markers"', html)
        self.assertEqual(html.count('name="question_image_file"'), 1)
        # Each blank row carries a real media slot for its hint (names assigned per row at submit).
        self.assertIn('data-media-base="blank_media"', html)
        self.assertNotIn('در حال توسعه', html.split('data-row-list="blank"')[1].split('</section>')[0])
        single = self.client.get(f'{url}?type=single').content.decode()
        self.assertNotIn('class="qf-drop qb-fb-add"', single)

    def test_true_false_question_saves_chosen_answer(self):
        self.client.force_login(self.teacher_user)
        response = self._create('true_false', {
            'text': 'سوال درست نادرست تست',
            'option_text': ['درست', 'نادرست'],
            'correct_options': ['2'],
        })
        self.assertEqual(response.status_code, 302)
        row = self._row('سوال درست نادرست تست')
        options = json.loads(row[3])
        self.assertEqual([item['text'] for item in options], ['درست', 'نادرست'])
        self.assertEqual(json.loads(row[4]), options[1]['id'])

    def test_question_bank_type_strip_filters_list(self):
        self.client.force_login(self.teacher_user)
        self._create('true_false', {'text': 'فیلتر نوع درست نادرست', 'option_text': ['درست', 'نادرست'], 'correct_options': ['1']})
        self._create('ordering', {'text': 'فیلتر نوع ترتیبی', 'ordering_item': ['الف', 'ب']})
        page = self.client.get(f"{reverse('core:teacher_questions')}?type=ordering&q=فیلتر نوع")
        html = page.content.decode()
        self.assertIn('فیلتر نوع ترتیبی', html)
        self.assertNotIn('فیلتر نوع درست نادرست', html)
        # کارت فعال اکنون باید فیلتر را بردارد و بقیه‌ی فیلترها حفظ شوند
        self.assertIn('class="is-active" aria-current="true"', html)
        self.assertIn('?q=%D9%81%DB%8C%D9%84%D8%AA%D8%B1%20%D9%86%D9%88%D8%B9&amp;difficulty=all&amp;status=all&amp;type=all', html)
        self.assertNotIn('teacher_question_create', html.split('question-type-strip')[1].split('</nav>')[0])


class AttemptGradingTests(_ScopedOrgFixture):
    """
    رگرسیون باگ‌های تصحیح خودکار: پاسخ و correct_answer هر دو JSON-بسته‌بندی
    ذخیره می‌شوند و قبلاً بدون باز کردن مقایسه می‌شدند؛ چندپاسخی فقط یک
    مقدار ذخیره می‌کرد؛ تطبیقی/مرتب‌سازی از ستون‌های اختصاصی خوانده نمی‌شدند.
    """

    def _score(self, question_columns, answer_value):
        from apps.core.views import grade_attempt
        question_id = str(uuid.uuid4())
        _insert(
            'questions',
            ['id', 'teacher_id', 'course_id', 'text', 'default_points'] + list(question_columns),
            [question_id, self.teacher_in_id, self.course_in_id, 'سوال تصحیح', 2] + list(question_columns.values()),
        )
        exam_id = str(uuid.uuid4())
        _insert('exams', ['id', 'title', 'teacher_id', 'course_id', 'start_at'],
                [exam_id, 'آزمون تصحیح', self.teacher_in_id, self.course_in_id, timezone.now()])
        _insert('exam_questions', ['id', 'exam_id', 'question_id', 'points', 'order_index'],
                [str(uuid.uuid4()), exam_id, question_id, 2, 1])
        attempt_id = str(uuid.uuid4())
        _insert('exam_attempts', ['id', 'exam_id', 'student_id', 'started_at', 'status'],
                [attempt_id, exam_id, self.student_in_id, timezone.now(), 'in_progress'])
        _insert('attempt_answers', ['id', 'attempt_id', 'question_id', 'answer', 'needs_manual_grading'],
                [str(uuid.uuid4()), attempt_id, question_id, json.dumps({'value': answer_value}, ensure_ascii=False), False])
        return grade_attempt(attempt_id)['score']

    def test_single_choice_graded_after_unwrapping_json(self):
        correct = {'type': 'single', 'correct_answer': json.dumps('2')}
        self.assertEqual(self._score(correct, '2'), 2)
        self.assertEqual(self._score(correct, '1'), 0)

    def test_multi_choice_compares_full_selection(self):
        correct = {'type': 'multi', 'correct_answer': json.dumps(['1', '2'])}
        self.assertEqual(self._score(correct, ['1', '2']), 2)
        self.assertEqual(self._score(correct, ['1']), 0)

    def test_fill_blank_ordering_and_matching(self):
        fill = {'type': 'fill_blank', 'correct_answer': json.dumps('الف|ب', ensure_ascii=False)}
        self.assertEqual(self._score(fill, 'الف|ب'), 2)
        self.assertEqual(self._score(fill, 'ب|الف'), 0)
        ordering = {'type': 'ordering', 'ordering_items': json.dumps([{'order': 1, 'text': 'x'}, {'order': 2, 'text': 'y'}, {'order': 3, 'text': 'z'}])}
        self.assertEqual(self._score(ordering, '1,2,3'), 2)
        self.assertEqual(self._score(ordering, '2,1,3'), 0)
        matching = {'type': 'matching', 'matching_pairs': json.dumps([{'left': 'a', 'right': 'b'}, {'left': 'c', 'right': 'd'}])}
        self.assertEqual(self._score(matching, '1:1,2:2'), 2)
        self.assertEqual(self._score(matching, '1:2,2:1'), 0)

    def test_negative_marking_applies_only_to_wrong_answers(self):
        single = {'type': 'single', 'correct_answer': json.dumps('2'),
                  'scoring_settings': json.dumps({'negative_marking': True, 'negative_value': 0.25})}
        self.assertEqual(self._score(single, '2'), 2)
        self.assertEqual(self._score(single, '1'), 0)  # total attempt score is floored at zero
        from apps.core.views import _grade_question
        row = {'type': 'single', 'correct_answer': json.dumps('2'), 'scoring_settings': {'negative_marking': True, 'negative_value': 0.25}}
        self.assertEqual(_grade_question(row, '1', 2)[0], -0.5)
        self.assertEqual(_grade_question(row, '', 2)[0], 0)

    def test_multi_partial_credit_and_selection_limit(self):
        from apps.core.views import _grade_question
        options = [{'id': str(i), 'text': t, 'points': '1', 'negative_points': '-0.5'} for i, t in enumerate('ابجد', start=1)]
        row = {'type': 'multi', 'correct_answer': json.dumps(['1', '2']), 'options': options,
               'scoring_settings': {'scoring_mode': 'partial', 'max_correct': 3, 'floor_zero': True}}
        self.assertEqual(_grade_question(row, ['1', '2'], 4)[0], 4)
        self.assertEqual(_grade_question(row, ['1'], 4)[0], 2)
        self.assertEqual(_grade_question(row, ['1', '3'], 4)[0], 1)  # (1 − 0.5) / 2 of the points
        self.assertEqual(_grade_question(row, ['3', '4'], 4)[0], 0)  # floored at zero
        self.assertEqual(_grade_question(row, ['1', '2', '3', '4'], 4)[0], 0)  # above max selection
        options[0]['points'] = '3'  # weighted option
        self.assertEqual(_grade_question(row, ['1'], 4)[0], 3)
        row['scoring_settings']['floor_zero'] = False
        self.assertEqual(_grade_question(row, ['3', '4'], 4)[0], -1)  # (0 − 1) / 4 of the points

    def test_short_answer_auto_grading_rules(self):
        from apps.core.views import _grade_question
        accepted = [
            {'answer': 'آتروپین', 'points': 100, 'match_mode': 'exact', 'case_sensitive': False, 'normalize': True, 'tolerance': 0},
            {'answer': '120', 'points': 100, 'match_mode': 'exact', 'case_sensitive': False, 'normalize': True, 'tolerance': 5},
            {'answer': 'adrenal', 'points': 50, 'match_mode': 'contains', 'case_sensitive': False, 'normalize': True, 'tolerance': 0},
        ]
        row = {'type': 'short_answer', 'accepted_answers': accepted}
        self.assertEqual(_grade_question(row, ' آتروپين ', 2), (2, False, True))  # Arabic ي normalised
        self.assertEqual(_grade_question(row, '۱۲۴', 2)[0], 2)  # Persian digits + tolerance
        self.assertEqual(_grade_question(row, '130', 2), (0, True, False))  # unmatched → manual review
        self.assertEqual(_grade_question(row, '', 2), (0, False, False))
        self.assertEqual(_grade_question(row, 'Adrenal gland', 2)[0], 1)
        legacy = {'type': 'short_answer', 'accepted_answers': [{'answer': 'x', 'points': '1'}]}
        self.assertEqual(_grade_question(legacy, 'x', 2), (0, True, False))

    def test_short_answer_typo_word_order_and_partial_switch(self):
        from apps.core.views import _grade_question
        accepted = [
            {'answer': 'سندرم کوشینگ', 'points': 100, 'match_mode': 'exact', 'normalize': True, 'typo_tolerance': 1, 'ignore_word_order': True},
            {'answer': 'cushing', 'points': 50, 'match_mode': 'exact', 'normalize': True, 'typo_tolerance': 0},
            {'answer': 'کم', 'points': 100, 'match_mode': 'exact', 'normalize': True, 'typo_tolerance': 2},
        ]
        row = {'type': 'short_answer', 'accepted_answers': accepted}
        self.assertEqual(_grade_question(row, 'کوشینگ سندرم', 2), (2, False, True))  # word order ignored
        self.assertEqual(_grade_question(row, 'سندروم کوشینگ', 2)[0], 2)  # one-letter typo forgiven
        self.assertEqual(_grade_question(row, 'سندروم کوشنگ', 2), (0, True, False))  # two typos > tolerance
        self.assertEqual(_grade_question(row, 'کو', 2), (0, True, False))  # short answers never fuzzy-matched
        self.assertEqual(_grade_question(row, 'Cushing', 2), (1, False, False))  # 50% < default 100% threshold
        row['scoring_settings'] = {'correct_threshold': 50}
        self.assertEqual(_grade_question(row, 'Cushing', 2), (1, False, True))
        row['scoring_settings'] = {'scoring_mode': 'all_or_nothing'}
        self.assertEqual(_grade_question(row, 'Cushing', 2), (0, False, False))  # matched, so no manual review
        self.assertEqual(_grade_question(row, 'سندرم کوشینگ', 2), (2, False, True))

    def test_fill_blank_alternatives_and_partial(self):
        from apps.core.views import _grade_question
        blanks = [
            {'blank': 1, 'answers': ['قلب', 'heart'], 'points': 1, 'case_sensitive': False, 'ignore_spaces': True},
            {'blank': 2, 'answers': ['ریه'], 'points': 3, 'case_sensitive': False, 'ignore_spaces': True},
        ]
        row = {'type': 'fill_blank', 'accepted_answers': blanks, 'scoring_settings': {'scoring_mode': 'partial'}}
        self.assertEqual(_grade_question(row, ['HEART', 'ریه'], 4)[0], 4)
        self.assertEqual(_grade_question(row, ['قلب', 'کبد'], 4)[0], 1)
        row['scoring_settings'] = {'scoring_mode': 'all_or_nothing'}
        self.assertEqual(_grade_question(row, ['قلب', 'کبد'], 4)[0], 0)

    def test_fill_blank_normalization_can_be_turned_off(self):
        from apps.core.views import _grade_question
        blanks = [{'blank': 1, 'answers': ['کلية ۱۲'], 'points': 1, 'case_sensitive': False, 'ignore_spaces': False}]
        row = {'type': 'fill_blank', 'accepted_answers': blanks, 'scoring_settings': {'scoring_mode': 'partial'}}
        # Default (and «نرمال‌سازی پاسخ‌ها» on): Arabic «ي/ك» and Latin digits match their Persian forms.
        self.assertEqual(_grade_question(row, ['کليه 12'], 2)[0], 2)
        row['scoring_settings'] = {'scoring_mode': 'partial', 'normalize_answers': False}
        self.assertEqual(_grade_question(row, ['کليه 12'], 2)[0], 0)
        self.assertEqual(_grade_question(row, ['کلية ۱۲'], 2)[0], 2)

    def test_ordering_and_matching_partial_modes(self):
        from apps.core.views import _grade_question
        items = [{'order': i, 'text': str(i)} for i in range(1, 5)]
        row = {'type': 'ordering', 'ordering_items': items, 'scoring_settings': {'scoring_mode': 'partial', 'ordering_mode': 'position'}}
        self.assertEqual(_grade_question(row, '1,2,4,3', 4)[0], 2)
        row['scoring_settings']['ordering_mode'] = 'adjacent'
        self.assertEqual(_grade_question(row, '2,3,4,1', 3)[0], 2)
        row['scoring_settings']['ordering_mode'] = 'exact'
        self.assertEqual(_grade_question(row, '1,2,4,3', 4)[0], 0)
        pairs = [{'left': 'a', 'right': 'x'}, {'left': 'b', 'right': 'x'}, {'left': 'c', 'right': 'y'}]
        matching = {'type': 'matching', 'matching_pairs': pairs, 'scoring_settings': {'scoring_mode': 'partial', 'matching_mode': 'many_to_one'}}
        self.assertEqual(_grade_question(matching, '1:1,2:1,3:3', 3)[0], 3)  # same right text accepted
        self.assertEqual(_grade_question(matching, '1:1,2:3,3:d1', 3)[0], 1)


class OsceBuilderTests(_ScopedOrgFixture):
    """سازنده‌ی آزمون OSCE: ذخیره‌ی واقعی ایستگاه‌ها و محدودشدن دسترسی به استاد مالک."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        User = get_user_model()
        cls.teacher_user = User.objects.create_user(username='teacher_in_scope', email='teacher_in_scope@example.com', password='x')
        cls.other_teacher_user = User.objects.create_user(username='teacher_out_scope', email='teacher_out_scope@example.com', password='x')

    def _basic(self, **overrides):
        data = {
            'tab': 'basic', 'exam_title': 'آزمون OSCE کودکان', 'course_id': self.course_in_id,
            'title': 'ارزیابی معاینه قلب', 'code': 'OSCE-01', 'station_type': 'clinical', 'domain': 'قلب و عروق کودکان',
            'difficulty': 'medium', 'description': 'ارزیابی مهارت معاینه قلب', 'position': '1', 'duration_minutes': '10',
            'is_active': '1', 'objectives': ['انجام معاینه دقیق قلب', 'ارتباط مؤثر با والدین'],
        }
        data.update(overrides)
        return data

    def _create_osce(self):
        self.client.force_login(self.teacher_user)
        response = self.client.post(reverse('core:teacher_osce_create'), self._basic())
        self.assertEqual(response.status_code, 302)
        with connection.cursor() as cursor:
            cursor.execute('SELECT id FROM osce_exams WHERE teacher_id = %s', [self.teacher_in_id])
            return str(cursor.fetchone()[0])

    def _stations(self, osce_id):
        with connection.cursor() as cursor:
            cursor.execute(
                'SELECT id, position, title, code, objectives, details, is_active FROM osce_stations WHERE osce_exam_id = %s ORDER BY position',
                [osce_id],
            )
            return cursor.fetchall()

    def test_structure_page_routes_osce_to_builder(self):
        self.client.force_login(self.teacher_user)
        response = self.client.get(reverse('core:teacher_question_type'))
        self.assertContains(response, reverse('core:teacher_osce_create'))

    def test_create_saves_exam_and_first_station(self):
        osce_id = self._create_osce()
        stations = self._stations(osce_id)
        self.assertEqual(len(stations), 1)
        self.assertEqual(stations[0][2], 'ارزیابی معاینه قلب')
        self.assertEqual(json.loads(stations[0][4]), ['انجام معاینه دقیق قلب', 'ارتباط مؤثر با والدین'])
        response = self.client.get(reverse('core:teacher_osce_builder', args=[osce_id]))
        self.assertContains(response, 'ارزیابی معاینه قلب')
        self.assertContains(response, 'tab=settings')

    def test_create_rejects_invalid_input_and_foreign_course(self):
        self.client.force_login(self.teacher_user)
        response = self.client.post(
            reverse('core:teacher_osce_create'),
            self._basic(title='', duration_minutes='0', course_id=self.course_out_id),
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'عنوان ایستگاه الزامی است.')
        self.assertContains(response, 'درس انتخاب‌شده معتبر نیست.')
        with connection.cursor() as cursor:
            cursor.execute('SELECT COUNT(*) FROM osce_exams')
            self.assertEqual(cursor.fetchone()[0], 0)

    def test_type_tab_saves_details(self):
        osce_id = self._create_osce()
        station_id = str(self._stations(osce_id)[0][0])
        response = self.client.post(reverse('core:teacher_osce_builder', args=[osce_id]), {
            'tab': 'type', 'station_id': station_id, 'station_type': 'skill', 'execution_mode': 'mannequin',
            'assessment_method': 'checklist', 'examiner_present': '1',
        })
        self.assertEqual(response.status_code, 302)
        details = json.loads(self._stations(osce_id)[0][5])
        self.assertEqual(details['execution_mode'], 'mannequin')
        self.assertTrue(details['examiner_present'])
        self.assertFalse(details['needs_standardized_patient'])

    def test_duplicate_code_is_rejected(self):
        osce_id = self._create_osce()
        self.client.post(reverse('core:teacher_osce_station_add', args=[osce_id]))
        second_id = str(self._stations(osce_id)[1][0])
        response = self.client.post(
            reverse('core:teacher_osce_builder', args=[osce_id]),
            self._basic(station_id=second_id, position='2'),
        )
        self.assertContains(response, 'این کد ایستگاه قبلاً در همین آزمون استفاده شده است.')

    def test_add_reorder_and_delete_stations(self):
        osce_id = self._create_osce()
        self.client.post(reverse('core:teacher_osce_station_add', args=[osce_id]))
        self.client.post(reverse('core:teacher_osce_station_add', args=[osce_id]))
        ids = [str(row[0]) for row in self._stations(osce_id)]
        self.assertEqual([row[3] for row in self._stations(osce_id)], ['OSCE-01', 'OSCE-02', 'OSCE-03'])
        url = reverse('core:teacher_osce_station_reorder', args=[osce_id])
        response = self.client.post(url, json.dumps({'order': list(reversed(ids))}), content_type='application/json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual([str(row[0]) for row in self._stations(osce_id)], list(reversed(ids)))
        bad = self.client.post(url, json.dumps({'order': ids[:2]}), content_type='application/json')
        self.assertEqual(bad.status_code, 400)
        self.client.post(reverse('core:teacher_osce_station_delete', args=[osce_id, ids[1]]))
        self.assertEqual([row[1] for row in self._stations(osce_id)], [1, 2])

    def test_other_teacher_cannot_read_or_change_osce(self):
        osce_id = self._create_osce()
        station_id = str(self._stations(osce_id)[0][0])
        self.client.force_login(self.other_teacher_user)
        self.assertEqual(self.client.get(reverse('core:teacher_osce_builder', args=[osce_id])).status_code, 404)
        self.assertEqual(
            self.client.post(reverse('core:teacher_osce_builder', args=[osce_id]), self._basic(station_id=station_id, title='تغییر')).status_code,
            404,
        )
        self.assertEqual(self.client.post(reverse('core:teacher_osce_station_add', args=[osce_id])).status_code, 404)
        self.assertEqual(self.client.post(reverse('core:teacher_osce_station_delete', args=[osce_id, station_id])).status_code, 404)
        reorder = self.client.post(
            reverse('core:teacher_osce_station_reorder', args=[osce_id]),
            json.dumps({'order': [station_id]}), content_type='application/json',
        )
        self.assertEqual(reorder.status_code, 404)
        self.assertEqual(self.client.post(reverse('core:teacher_osce_settings', args=[osce_id]), {'exam_title': 'تغییر'}).status_code, 404)
        self.assertNotContains(self.client.get(reverse('core:teacher_osce_list')), 'آزمون OSCE کودکان')
        self.assertEqual(self._stations(osce_id)[0][2], 'ارزیابی معاینه قلب')

    def test_station_id_from_another_exam_is_not_editable(self):
        osce_id = self._create_osce()
        self.client.force_login(self.other_teacher_user)
        other = self.client.post(reverse('core:teacher_osce_create'), self._basic(course_id='', code='X-1'))
        self.assertEqual(other.status_code, 302)
        with connection.cursor() as cursor:
            cursor.execute(
                'SELECT s.id FROM osce_stations s JOIN osce_exams o ON o.id = s.osce_exam_id WHERE o.teacher_id = %s',
                [self.teacher_out_id],
            )
            foreign_station = str(cursor.fetchone()[0])
        self.client.force_login(self.teacher_user)
        response = self.client.post(reverse('core:teacher_osce_builder', args=[osce_id]), self._basic(station_id=foreign_station))
        self.assertEqual(response.status_code, 404)

    def _details(self, osce_id):
        return json.loads(self._stations(osce_id)[0][5])

    def _post_tab(self, osce_id, tab, data):
        station_id = str(self._stations(osce_id)[0][0])
        return self.client.post(reverse('core:teacher_osce_builder', args=[osce_id]), {'tab': tab, 'station_id': station_id, **data})

    def test_scenario_tab_saves_rich_text_media_and_new_objective(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        osce_id = self._create_osce()
        story = {'ops': [{'insert': 'کودک ۷ ساله '}, {'insert': 'VSD', 'attributes': {'bold': True}}, {'insert': '\n'}]}
        with self.settings(MEDIA_ROOT=tempfile.mkdtemp()):
            response = self._post_tab(osce_id, 'scenario', {
                'patient_age': '۷ سال', 'patient_gender': 'boy', 'patient_chief_complaint': 'تنگی نفس', 'patient_role': 'parent',
                'scenario_story': json.dumps(story), 'scenario_notes': 'نکته ایمنی',
                'focus_objectives': ['انجام معاینه دقیق قلب'], 'new_objective': 'تصمیم‌گیری بالینی اولیه',
                'scenario_media': [
                    SimpleUploadedFile('cxr.png', b'\x89PNG\r\n\x1a\n' + b'0' * 32, content_type='image/png'),
                    SimpleUploadedFile('lab.pdf', b'%PDF-1.4 test', content_type='application/pdf'),
                ],
            })
            self.assertEqual(response.status_code, 302)
            scenario = self._details(osce_id)['scenario']
            self.assertEqual(scenario['patient']['role'], 'parent')
            self.assertEqual(scenario['story'], story)
            self.assertEqual([item['kind'] for item in scenario['media']], ['image', 'pdf'])
            self.assertIn('تصمیم‌گیری بالینی اولیه', scenario['focus_objectives'])
            self.assertIn('تصمیم‌گیری بالینی اولیه', json.loads(self._stations(osce_id)[0][4]))
            removed = self._post_tab(osce_id, 'scenario', {'remove_media': [scenario['media'][0]['url']]})
            self.assertEqual(removed.status_code, 302)
            self.assertEqual([item['name'] for item in self._details(osce_id)['scenario']['media']], ['lab.pdf'])

    def test_scenario_upload_rejects_disguised_file_and_keeps_old_data(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        osce_id = self._create_osce()
        with self.settings(MEDIA_ROOT=tempfile.mkdtemp()):
            response = self._post_tab(osce_id, 'scenario', {
                'scenario_media': [SimpleUploadedFile('report.pdf', b'<script>alert(1)</script>', content_type='application/pdf')],
            })
            self.assertEqual(response.status_code, 200)
            self.assertContains(response, 'با پسوند آن هم‌خوانی ندارد')
            response = self._post_tab(osce_id, 'scenario', {
                'scenario_media': [SimpleUploadedFile('run.exe', b'MZ', content_type='application/octet-stream')],
            })
            self.assertContains(response, 'مجاز نیست')
        self.assertNotIn('scenario', self._details(osce_id))

    def test_scenario_requires_patient_when_station_needs_standardized_patient(self):
        osce_id = self._create_osce()
        self._post_tab(osce_id, 'type', {'station_type': 'clinical', 'execution_mode': 'standardized_patient',
                                         'assessment_method': 'checklist', 'needs_standardized_patient': '1'})
        response = self._post_tab(osce_id, 'scenario', {'patient_age': '۷ سال'})
        self.assertContains(response, 'این ایستگاه به بیمار استانداردشده نیاز دارد')

    def test_instructions_tab_requires_body_and_renders_escaped_preview(self):
        osce_id = self._create_osce()
        response = self._post_tab(osce_id, 'instructions', {'instructions_body': json.dumps({'ops': [{'insert': '\n'}]})})
        self.assertContains(response, 'متن دستورالعمل دانشجو الزامی است.')
        body = {'ops': [{'insert': '<script>alert(1)</script> معاینه کنید'}, {'insert': '\n', 'attributes': {'list': 'ordered'}}]}
        response = self._post_tab(osce_id, 'instructions', {'instructions_body': json.dumps(body), 'show_task_list': '1'})
        self.assertEqual(response.status_code, 302)
        instructions = self._details(osce_id)['instructions']
        self.assertTrue(instructions['show_task_list'])
        self.assertFalse(instructions['show_start_button'])
        page = self.client.get(response['Location'])
        self.assertContains(page, '<ol><li>&lt;script&gt;alert(1)&lt;/script&gt; معاینه کنید</li></ol>', html=False)
        self.assertNotContains(page, '<script>alert(1)</script>')

    def _criteria_post(self, **overrides):
        data = {
            'criteria_uid': ['criteria-0', 'criteria-1', 'criteria-2'],
            'criteria_text': ['معرفی خود', 'معاینه قلب به ترتیب صحیح', ''],
            'criteria_type': ['binary', 'scale', 'binary'],
            'criteria_points': ['2', '8.5', '1'],
            'criteria_guide': ['خود را معرفی می‌کند', '', ''],
            'criteria_required': ['criteria-1'],
            'global_uid': ['global-0'], 'global_text': ['عملکرد کلی'], 'global_points': ['4'], 'global_guide': [''],
            'level_label': ['عالی', 'خوب', 'متوسط', 'ضعیف تا متوسط', 'ضعیف'],
            'level_description': ['کامل', 'جزئی', 'با کمک', 'ناقص', 'انجام نمی‌دهد'],
            'scoring_mode': 'per_item', 'show_descriptions': '1',
        }
        data.update(overrides)
        return data

    def test_criteria_tab_saves_items_levels_and_totals(self):
        osce_id = self._create_osce()
        response = self._post_tab(osce_id, 'criteria', self._criteria_post())
        self.assertEqual(response.status_code, 302)
        criteria = self._details(osce_id)['criteria']
        self.assertEqual([item['text'] for item in criteria['items']], ['معرفی خود', 'معاینه قلب به ترتیب صحیح'])
        self.assertEqual([item['required'] for item in criteria['items']], [False, True])
        self.assertEqual(criteria['items'][1]['points'], '8.5')
        self.assertEqual(criteria['total_points'], '14.5')
        self.assertEqual(criteria['global_items'][0]['type'], 'scale')
        self.assertEqual(criteria['scale_levels'][3]['label'], 'ضعیف تا متوسط')
        self.assertEqual(criteria['display'], {'show_descriptions': True, 'show_item_points': False,
                                               'show_running_total': False, 'show_examiner_notes': False})
        page = self.client.get(response['Location'])
        self.assertContains(page, 'معاینه قلب به ترتیب صحیح')
        self.assertContains(page, 'معاینه قلب کودکان')  # built-in template list

    def test_criteria_tab_validates_points_pass_score_and_method(self):
        osce_id = self._create_osce()
        response = self._post_tab(osce_id, 'criteria', self._criteria_post(
            criteria_points=['0', 'abc', '1'], scoring_mode='pass_fail', pass_score='99'))
        self.assertContains(response, 'امتیاز باید عددی بزرگ‌تر از صفر')
        self.assertContains(response, 'حد نصاب قبولی باید')
        self._post_tab(osce_id, 'type', {'station_type': 'clinical', 'execution_mode': 'mannequin', 'assessment_method': 'combined'})
        response = self._post_tab(osce_id, 'criteria', self._criteria_post(global_uid=[], global_text=[], global_points=[], global_guide=[]))
        self.assertContains(response, 'حداقل یک مورد ارزیابی کلی')
        self.assertNotIn('criteria', self._details(osce_id))

    def test_teacher_templates_are_saved_listed_and_scoped(self):
        url = reverse('core:teacher_osce_template_save')
        self.client.force_login(self.teacher_user)
        items = [{'text': 'شستن دست', 'type': 'binary', 'points': '1', 'required': True, 'guide': ''}]
        response = self.client.post(url, json.dumps({'title': 'قالب من', 'items': items}), content_type='application/json')
        self.assertEqual(response.status_code, 200)
        template_id = response.json()['template']['id']
        duplicate = self.client.post(url, json.dumps({'title': 'قالب من', 'items': items}), content_type='application/json')
        self.assertEqual(duplicate.status_code, 400)
        invalid = self.client.post(url, json.dumps({'title': 'بد', 'items': [{'text': 'x', 'type': 'evil', 'points': '-1'}]}), content_type='application/json')
        self.assertEqual(invalid.status_code, 400)
        osce_id = self._create_osce()
        station_id = str(self._stations(osce_id)[0][0])
        page = self.client.get(reverse('core:teacher_osce_builder', args=[osce_id]) + f'?station={station_id}&tab=criteria')
        self.assertContains(page, 'قالب من')
        delete_url = reverse('core:teacher_osce_template_delete', args=[template_id])
        self.client.force_login(self.other_teacher_user)
        self.assertEqual(self.client.post(delete_url).status_code, 404)
        self.client.force_login(self.teacher_user)
        self.assertEqual(self.client.post(delete_url).status_code, 200)
        with connection.cursor() as cursor:
            cursor.execute('SELECT COUNT(*) FROM osce_criteria_templates')
            self.assertEqual(cursor.fetchone()[0], 0)

    def _timing_post(self, **overrides):
        data = {
            'duration_minutes': '8', 'study_minutes': '2', 'transition_minutes': '1', 'execution_pattern': 'rotation',
            'capacity': '1', 'groups_count': '4', 'students_per_group': '6', 'start_time': '09:00',
            'primary_examiner': self.teacher_in_id, 'second_examiner': '', 'alert_type': 'audio_visual',
            'timer_student': '1', 'end_sound': '1',
        }
        data.update(overrides)
        return data

    def _settings_post(self, **overrides):
        data = {'access_scope': 'selected', 'max_attempts': '1', 'exam_pass_score': '60', 'negative_mark': '0.25',
                'show_remaining_time': '1', 'notify_students': '1', 'tags': ['بالینی', 'ارتباطی', 'بالینی'], 'exam_notes': ''}
        data.update(overrides)
        return data

    def test_resources_tab_saves_tables_and_rejects_bad_rows(self):
        osce_id = self._create_osce()
        response = self._post_tab(osce_id, 'resources', {
            'equipment_name': ['تخت معاینه', ''], 'equipment_count': ['1', '1'], 'equipment_status': ['required', 'optional'],
            'supplies_name': ['دستکش'], 'supplies_count': ['10'], 'supplies_status': ['required'],
            'staff_name': ['بیمار استانداردشده'], 'staff_count': ['1'], 'staff_note': ['آموزش‌دیده'],
            'safety_notes': 'بررسی دستگاه قبل از شروع',
        })
        self.assertEqual(response.status_code, 302)
        resources = self._details(osce_id)['resources']
        self.assertEqual(resources['equipment'], [{'name': 'تخت معاینه', 'count': 1, 'status': 'required'}])
        self.assertEqual(resources['staff'][0]['note'], 'آموزش‌دیده')
        bad = self._post_tab(osce_id, 'resources', {
            'equipment_name': ['تخت'], 'equipment_count': ['0'], 'equipment_status': ['evil'],
            'supplies_name': ['a', 'b'], 'supplies_count': ['1'], 'supplies_status': ['required'],
        })
        self.assertContains(bad, 'تعداد باید عددی بین')
        self.assertContains(bad, 'وضعیت نامعتبر است')
        self.assertContains(bad, 'ناقص ارسال شده است')
        self.assertEqual(self._details(osce_id)['resources']['equipment'][0]['name'], 'تخت معاینه')

    def test_timing_tab_saves_and_limits_examiners_to_org_unit(self):
        osce_id = self._create_osce()
        outsider = self._post_tab(osce_id, 'timing', self._timing_post(primary_examiner=self.teacher_out_id))
        self.assertContains(outsider, 'ارزیاب اصلی را از فهرست انتخاب کنید.')
        response = self._post_tab(osce_id, 'timing', self._timing_post())
        self.assertEqual(response.status_code, 302)
        timing = self._details(osce_id)['timing']
        self.assertEqual(timing['study_minutes'], 2)
        self.assertTrue(timing['alerts']['timer_student'])
        self.assertFalse(timing['alerts']['one_minute_warning'])
        with connection.cursor() as cursor:
            cursor.execute('SELECT duration_minutes FROM osce_stations WHERE osce_exam_id = %s', [osce_id])
            self.assertEqual(cursor.fetchone()[0], 8)
            cursor.execute('SELECT settings FROM osce_exams WHERE id = %s', [osce_id])
            self.assertEqual(json.loads(cursor.fetchone()[0])['rotation']['start_time'], '09:00')
        page = self.client.get(response['Location'])
        self.assertContains(page, '۰۹:۰۰ - ۰۹:۱۰')
        self.assertContains(page, self.teacher_in_id)
        self.assertNotContains(page, self.teacher_out_id)

    def test_rotation_schedule_shifts_groups_and_runs_extra_circuits(self):
        from apps.core.osce_views import build_rotation_schedule
        stations = [
            {'title': f'S{i}', 'is_active': True, 'duration_minutes': 8,
             'details': {'timing': {'study_minutes': 2 if i == 1 else 0, 'transition_minutes': 1}}}
            for i in range(1, 4)
        ] + [{'title': 'off', 'is_active': False, 'duration_minutes': 5, 'details': {}}]
        schedule = build_rotation_schedule(stations, {'start_time': '09:00', 'groups_count': 4})
        self.assertEqual(schedule['stations'], ['S1', 'S2', 'S3'])
        self.assertEqual(schedule['round_minutes'], 10)
        self.assertEqual(schedule['rows'][0]['cells'], ['09:00 - 09:10', '09:11 - 09:21', '09:22 - 09:32'])
        self.assertEqual(schedule['rows'][1]['cells'], ['09:22 - 09:32', '09:00 - 09:10', '09:11 - 09:21'])
        self.assertEqual(schedule['rows'][3]['circuit'], 2)
        self.assertEqual(schedule['rows'][3]['cells'][0], '09:33 - 09:43')
        self.assertEqual(schedule['end_time'], '10:05')
        self.assertIsNone(build_rotation_schedule(stations, {'start_time': '25:00', 'groups_count': 1}))

    def test_settings_tab_saves_exam_level_settings(self):
        osce_id = self._create_osce()
        bad = self._post_tab(osce_id, 'settings', self._settings_post(negative_mark='2', max_attempts='9', access_scope='all'))
        self.assertContains(bad, 'نمره منفی باید عددی بین ۰ و ۱ باشد.')
        self.assertContains(bad, 'محدودیت تعداد تلاش')
        response = self._post_tab(osce_id, 'settings', self._settings_post())
        self.assertEqual(response.status_code, 302)
        with connection.cursor() as cursor:
            cursor.execute('SELECT settings FROM osce_exams WHERE id = %s', [osce_id])
            exam = json.loads(cursor.fetchone()[0])['exam']
        self.assertEqual(exam['tags'], ['بالینی', 'ارتباطی'])
        self.assertTrue(exam['show_remaining_time'])
        self.assertFalse(exam['show_progress'])

    def test_publish_requires_complete_exam_and_locks_editing(self):
        osce_id = self._create_osce()
        preview_url = reverse('core:teacher_osce_preview', args=[osce_id])
        publish_url = reverse('core:teacher_osce_publish', args=[osce_id])
        self.assertContains(self.client.get(preview_url), 'دستورالعمل دانشجو ثبت نشده است')
        self.client.post(publish_url, {'action': 'publish'})
        with connection.cursor() as cursor:
            cursor.execute('SELECT status FROM osce_exams WHERE id = %s', [osce_id])
            self.assertEqual(cursor.fetchone()[0], 'draft')
        self._post_tab(osce_id, 'type', {'station_type': 'clinical', 'execution_mode': 'mannequin', 'assessment_method': 'checklist'})
        self._post_tab(osce_id, 'instructions', {'instructions_body': json.dumps({'ops': [{'insert': 'معاینه کنید\n'}]})})
        self._post_tab(osce_id, 'criteria', self._criteria_post(global_uid=[], global_text=[], global_points=[], global_guide=[]))
        self._post_tab(osce_id, 'timing', self._timing_post())
        self._post_tab(osce_id, 'settings', self._settings_post())
        self.assertContains(self.client.get(preview_url), 'آماده انتشار است')
        self.client.post(publish_url, {'action': 'publish'})
        with connection.cursor() as cursor:
            cursor.execute('SELECT status FROM osce_exams WHERE id = %s', [osce_id])
            self.assertEqual(cursor.fetchone()[0], 'published')
        self._post_tab(osce_id, 'basic', self._basic(title='تغییر بعد از انتشار'))
        self.assertEqual(self._stations(osce_id)[0][2], 'ارزیابی معاینه قلب')
        self.client.post(reverse('core:teacher_osce_station_add', args=[osce_id]))
        self.assertEqual(len(self._stations(osce_id)), 1)
        self.client.force_login(self.other_teacher_user)
        self.assertEqual(self.client.get(preview_url).status_code, 404)
        self.assertEqual(self.client.post(publish_url, {'action': 'unpublish'}).status_code, 404)
        self.client.force_login(self.teacher_user)
        self.client.post(publish_url, {'action': 'unpublish'})
        self._post_tab(osce_id, 'basic', self._basic(title='ویرایش پس از برگشت'))
        self.assertEqual(self._stations(osce_id)[0][2], 'ویرایش پس از برگشت')

    def test_legacy_osce_builder_redirects_to_new_builder(self):
        self.client.force_login(self.teacher_user)
        new_url = reverse('core:teacher_osce_create')
        self.assertRedirects(self.client.get(reverse('core:teacher_question_create') + '?type=osce'), new_url, fetch_redirect_response=False)
        response = self.client.post(reverse('core:teacher_question_create'), {'type': 'osce', 'text': 'ایستگاه قدیمی'})
        self.assertRedirects(response, new_url, fetch_redirect_response=False)
        with connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) FROM questions WHERE text = 'ایستگاه قدیمی'")
            self.assertEqual(cursor.fetchone()[0], 0)
        self.assertContains(self.client.get(reverse('core:teacher_exams')), reverse('core:teacher_osce_list'))

    def test_template_library_is_shared_within_org_unit_only(self):
        colleague_id = self._make_profile('teacher_colleague', 'استاد همکار')
        _insert('teacher_profiles', ['user_id', 'org_unit_id'], [colleague_id, self.ou_scoped])
        _insert('user_roles', ['id', 'user_id', 'role'], [str(uuid.uuid4()), colleague_id, 'teacher'])
        colleague = get_user_model().objects.create_user(username='teacher_colleague', email='teacher_colleague@example.com', password='x')
        items = [{'text': 'معاینه ریه', 'type': 'binary', 'points': '2', 'required': False, 'guide': ''}]
        self.client.force_login(self.teacher_user)
        saved = self.client.post(reverse('core:teacher_osce_template_save'), json.dumps({'title': 'قالب ریه', 'items': items}), content_type='application/json')
        template = saved.json()['template']
        library_url = reverse('core:teacher_osce_template_library')

        self.client.force_login(colleague)
        self.assertEqual(self.client.get(library_url).json()['templates'], [])  # not shared yet
        self.assertEqual(self.client.post(template['share_url'], json.dumps({'shared': True}), content_type='application/json').status_code, 404)

        self.client.force_login(self.teacher_user)
        self.assertTrue(self.client.post(template['share_url'], json.dumps({'shared': True}), content_type='application/json').json()['shared'])
        self.assertEqual(self.client.get(library_url).json()['templates'], [])  # own templates are not listed

        self.client.force_login(colleague)
        shared = self.client.get(library_url).json()['templates']
        self.assertEqual([(item['title'], item['author']) for item in shared], [('قالب ریه', 'استاد داخل محدوده')])
        self.assertEqual(self.client.get(library_url, {'q': 'ریه'}).json()['templates'][0]['items'][0]['text'], 'معاینه ریه')
        self.assertEqual(self.client.get(library_url, {'q': 'قلب'}).json()['templates'], [])

        self.client.force_login(self.other_teacher_user)  # other org unit
        self.assertEqual(self.client.get(library_url).json()['templates'], [])


class RichTextTests(TestCase):
    def test_delta_to_html_escapes_text_and_filters_attributes(self):
        from apps.core.rich_text import clean_delta, delta_to_html
        delta = {'ops': [
            {'insert': 'سلام <b>', 'attributes': {'bold': True, 'color': 'red', 'script': 'x'}},
            {'insert': 'لینک', 'attributes': {'link': 'javascript:alert(1)'}},
            {'insert': 'خوب', 'attributes': {'link': 'https://example.com/?a="b"'}},
            {'insert': {'image': 'https://evil.example/x.png'}},
            {'insert': '\n', 'attributes': {'header': 1, 'align': 'center'}},
            {'insert': 'یک'}, {'insert': '\n', 'attributes': {'list': 'bullet'}},
            {'insert': 'دو'}, {'insert': '\n', 'attributes': {'list': 'bullet'}},
            {'insert': 'پایان\n'},
        ]}
        html = delta_to_html(delta)
        self.assertEqual(
            html,
            '<h3 style="text-align:center"><strong>سلام &lt;b&gt;</strong>لینک'
            '<a href="https://example.com/?a=&quot;b&quot;" target="_blank" rel="noopener noreferrer">خوب</a></h3>'
            '<ul><li>یک</li><li>دو</li></ul><p>پایان</p>',
        )
        self.assertEqual(clean_delta('not json'), {'ops': []})
        self.assertEqual(clean_delta({'ops': 'x'}), {'ops': []})


class EssayQuestionBuilderTests(_ScopedOrgFixture):
    """سؤال تشریحی: صورت سؤال قالب‌دار (Quill)، دستورالعمل دانشجو، پاسخ نمونه و نمایش برای مصحح."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.teacher_user = get_user_model().objects.create_user(username='teacher_in_scope', email='t_in@example.com', password='x')

    def _post_essay(self, **overrides):
        data = {
            'type': 'essay', 'structure': 'independent', 'course_id': self.course_in_id, 'is_published': '1',
            'text': '',
            'text_delta': json.dumps({'ops': [
                {'insert': 'نقش '}, {'insert': 'بارورفلکس', 'attributes': {'bold': True}}, {'insert': ' را توضیح دهید <b>\n'},
            ]}),
            'student_instructions': 'با استناد به منابع درسی پاسخ دهید.',
            'sample_answer_text': 'بارورفلکس با حس فشار...',
            'rubric_criterion': ['دقت علمی'], 'rubric_weak': ['نادرست'], 'rubric_average': ['ناقص'],
            'rubric_excellent': ['دقیق'], 'rubric_points': ['10'], 'default_points': '10',
            'answer_formats': ['text', 'file'],
        }
        data.update(overrides)
        return self.client.post(f"{reverse('core:teacher_question_create')}?type=essay", data)

    def _essay_row(self):
        with connection.cursor() as cursor:
            cursor.execute("SELECT id, text, scoring_settings FROM questions WHERE type = 'essay' AND teacher_id = %s", [self.teacher_in_id])
            row = cursor.fetchone()
        return row and (str(row[0]), row[1], json.loads(row[2]))

    def test_essay_saves_rich_stem_instructions_and_sample_answer(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.force_login(self.teacher_user)
        with self.settings(MEDIA_ROOT=tempfile.mkdtemp()):
            response = self._post_essay(sample_answer_file=SimpleUploadedFile('sample.pdf', b'%PDF-1.4 x', content_type='application/pdf'))
            self.assertEqual(response.status_code, 302)
            question_id, text, settings_data = self._essay_row()
            self.assertEqual(text, 'نقش بارورفلکس را توضیح دهید <b>')
            essay = settings_data['essay']
            self.assertEqual(settings_data['text_delta']['ops'][1]['attributes'], {'bold': True})
            self.assertEqual(essay['student_instructions'], 'با استناد به منابع درسی پاسخ دهید.')
            self.assertTrue(essay['sample_answer_file']['url'].startswith('/media/question-media/'))
            self.assertEqual(essay['answer_formats'], ['text', 'file'])
            page = self.client.get(reverse('core:teacher_question_edit', args=[question_id]))
            self.assertContains(page, 'با استناد به منابع درسی پاسخ دهید.')
            self.assertContains(page, 'sample.pdf')
            self.assertContains(page, 'id="question-text-delta"')

    def test_essay_rejects_bad_sample_file_and_long_instructions(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.force_login(self.teacher_user)
        with self.settings(MEDIA_ROOT=tempfile.mkdtemp()):
            response = self._post_essay(
                student_instructions='x' * 1001,
                sample_answer_file=SimpleUploadedFile('run.exe', b'MZ', content_type='application/octet-stream'),
            )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'دستورالعمل دانشجو حداکثر')
        self.assertContains(response, 'مجاز نیست')
        self.assertIsNone(self._essay_row())

    def test_grading_aid_and_student_view_render_escaped_rich_text(self):
        from types import SimpleNamespace
        from apps.core.views import _essay_view, _grading_aid
        self.client.force_login(self.teacher_user)
        self._post_essay()
        question_id, _text, settings_data = self._essay_row()
        student_view = _essay_view(settings_data)
        self.assertEqual(student_view['text_html'], '<p>نقش <strong>بارورفلکس</strong> را توضیح دهید &lt;b&gt;</p>')
        aid = _grading_aid(SimpleNamespace(legacy_erd_id=question_id))
        self.assertEqual(aid['rubric'][0]['criterion'], 'دقت علمی')
        self.assertEqual(aid['rubric_total'], '10')
        self.assertEqual(aid['sample_answer_text'], 'بارورفلکس با حس فشار...')
        self.assertIsNone(_grading_aid(SimpleNamespace(legacy_erd_id='')))

    def test_student_attempt_page_shows_points_time_and_no_navigation(self):
        question_id, attempt_id = str(uuid.uuid4()), str(uuid.uuid4())
        _insert('questions', ['id', 'teacher_id', 'text', 'type', 'default_points', 'suggested_time_seconds'], [question_id, self.teacher_in_id, 'تشخیص را بنویسید', 'short_answer', 1, 120])
        _insert('exam_questions', ['id', 'exam_id', 'question_id', 'points', 'order_index'], [str(uuid.uuid4()), self.exam_in_id, question_id, 2.5, 1])
        _insert('exam_attempts', ['id', 'exam_id', 'student_id', 'started_at', 'status'], [attempt_id, self.exam_in_id, self.student_in_id, timezone.now(), 'in_progress'])
        student = get_user_model().objects.create_user(username='student_in_scope', email='student_in_scope@example.com', password='x')
        self.client.force_login(student)
        page = self.client.get(reverse('core:student_attempt', args=[attempt_id])).content.decode()
        self.assertIn('۲.۵ نمره', page)  # exam_questions.points wins over the question's default
        self.assertIn('زمان پیشنهادی: ۲ دقیقه', page)
        self.assertIn('سؤال ۱ — پاسخ کوتاه', page)
        self.assertIn('data-connection', page)
        # No way out of the attempt from the header: no dashboard link.
        self.assertNotIn(f'href="{reverse("core:dashboard")}"', page)
        # The mark checkbox sits in the side card but still submits with the answer form.
        self.assertIn(f'name="mark_{question_id}" form="answer-form"', page)

    def test_finish_button_saves_the_answer_before_leaving(self):
        question_id, attempt_id = str(uuid.uuid4()), str(uuid.uuid4())
        _insert('questions', ['id', 'teacher_id', 'text', 'type', 'default_points'], [question_id, self.teacher_in_id, 'تشخیص را بنویسید', 'short_answer', 1])
        _insert('exam_questions', ['id', 'exam_id', 'question_id', 'points', 'order_index'], [str(uuid.uuid4()), self.exam_in_id, question_id, 1, 1])
        _insert('exam_attempts', ['id', 'exam_id', 'student_id', 'started_at', 'status'], [attempt_id, self.exam_in_id, self.student_in_id, timezone.now(), 'in_progress'])
        self.client.force_login(get_user_model().objects.create_user(username='student_in_scope', email='student_in_scope@example.com', password='x'))
        url = reverse('core:student_attempt', args=[attempt_id])
        self.assertContains(self.client.get(url), 'form="answer-form" name="action" value="finish"')
        response = self.client.post(url, {f'answer_{question_id}': 'گاستریت', 'action': 'finish'})
        self.assertRedirects(response, reverse('core:student_attempt_submit', args=[attempt_id]), fetch_redirect_response=False)
        with connection.cursor() as cursor:
            cursor.execute('SELECT answer FROM attempt_answers WHERE attempt_id = %s AND question_id = %s', [attempt_id, question_id])
            saved = cursor.fetchone()
        self.assertIsNotNone(saved)
        self.assertIn('گاستریت', saved[0] if isinstance(saved[0], str) else json.dumps(saved[0], ensure_ascii=False))

    def test_student_attempt_page_renders_fill_blank_answers_and_essay(self):
        """Regression: the essay block must not swallow the fill-blank answer refill."""
        fill_id, essay_id, attempt_id = str(uuid.uuid4()), str(uuid.uuid4()), str(uuid.uuid4())
        _insert('questions', ['id', 'teacher_id', 'text', 'type', 'default_points'], [fill_id, self.teacher_in_id, 'قلب [۱] حفره دارد', 'fill_blank', 1])
        _insert('questions', ['id', 'teacher_id', 'text', 'type', 'default_points', 'scoring_settings'], [
            essay_id, self.teacher_in_id, 'نقش بارورفلکس', 'essay', 5,
            json.dumps({'essay': {'text_delta': {'ops': [{'insert': 'نقش '}, {'insert': 'بارورفلکس', 'attributes': {'bold': True}}, {'insert': '\n'}]},
                                  'student_instructions': 'با منابع پاسخ دهید.'}}),
        ])
        _insert('exam_questions', ['id', 'exam_id', 'question_id', 'points', 'order_index'], [str(uuid.uuid4()), self.exam_in_id, fill_id, 1, 1])
        _insert('exam_questions', ['id', 'exam_id', 'question_id', 'points', 'order_index'], [str(uuid.uuid4()), self.exam_in_id, essay_id, 5, 2])
        _insert('exam_attempts', ['id', 'exam_id', 'student_id', 'started_at', 'status'], [attempt_id, self.exam_in_id, self.student_in_id, timezone.now(), 'in_progress'])
        _insert('attempt_answers', ['id', 'attempt_id', 'question_id', 'answer'], [str(uuid.uuid4()), attempt_id, fill_id, json.dumps({'value': ['چهار']})])
        student = get_user_model().objects.create_user(username='student_in_scope', email='student_in_scope@example.com', password='x')
        self.client.force_login(student)
        url = reverse('core:student_attempt', args=[attempt_id])
        fill_page = self.client.get(url + '?q=1')
        self.assertEqual(fill_page.status_code, 200)
        self.assertContains(fill_page, 'value="چهار"')
        essay_page = self.client.get(url + '?q=2')
        self.assertEqual(essay_page.status_code, 200)
        self.assertContains(essay_page, '<strong>بارورفلکس</strong>', html=False)
        self.assertContains(essay_page, 'با منابع پاسخ دهید.')

    def test_student_attempt_renders_rich_fill_blank_and_media_hints(self):
        fill_id, attempt_id = str(uuid.uuid4()), str(uuid.uuid4())
        hint = '/media/question-media/skull.png'
        _insert('questions', ['id', 'teacher_id', 'text', 'type', 'default_points', 'scoring_settings', 'accepted_answers'], [
            fill_id, self.teacher_in_id, 'جمجمه [۱] و [۲]', 'fill_blank', 1,
            json.dumps({'text_delta': {'ops': [{'insert': 'جمجمه', 'attributes': {'bold': True}}, {'insert': ' [۱] و [۲]\n'}]}}),
            json.dumps([{'blank': 1, 'answers': ['نوروکرانیوم'], 'media': ''}, {'blank': 2, 'answers': ['ویسروکرانیوم'], 'media': hint}]),
        ])
        _insert('exam_questions', ['id', 'exam_id', 'question_id', 'points', 'order_index'], [str(uuid.uuid4()), self.exam_in_id, fill_id, 1, 1])
        _insert('exam_attempts', ['id', 'exam_id', 'student_id', 'started_at', 'status'], [attempt_id, self.exam_in_id, self.student_in_id, timezone.now(), 'in_progress'])
        _insert('attempt_answers', ['id', 'attempt_id', 'question_id', 'answer'], [str(uuid.uuid4()), attempt_id, fill_id, json.dumps({'value': ['پاسخ', '']})])
        student = get_user_model().objects.create_user(username='student_in_scope', email='student_in_scope@example.com', password='x')
        self.client.force_login(student)
        page = self.client.get(reverse('core:student_attempt', args=[attempt_id]) + '?q=1').content.decode()
        self.assertIn('qx-fillblank--rich', page)
        self.assertIn('<strong>جمجمه</strong>', page)
        self.assertEqual(page.count(f'name="answer_{fill_id}"'), 2)
        self.assertIn('value="پاسخ"', page)
        self.assertIn(f'<img src="{hint}"', page)
        # The hint list is built from accepted_answers, but the accepted answers themselves never reach the page.
        self.assertNotIn('نوروکرانیوم', page)
        self.assertNotIn('ویسروکرانیوم', page)

    def test_essay_requires_an_answer_format(self):
        self.client.force_login(self.teacher_user)
        response = self._post_essay(answer_formats=[])
        self.assertContains(response, 'حداقل یک قالب پاسخ')
        self.assertIsNone(self._essay_row())

    def test_rich_stem_is_saved_and_seeded_for_all_independent_types(self):
        from apps.core.views import _question_form_data
        self.client.force_login(self.teacher_user)
        delta = {'ops': [{'insert': 'کدام '}, {'insert': 'صحیح', 'attributes': {'italic': True}}, {'insert': ' است؟'}, {'insert': '\n'}]}
        response = self.client.post(f"{reverse('core:teacher_question_create')}?type=single", {
            'type': 'single', 'structure': 'independent', 'course_id': self.course_in_id, 'text': '', 'text_delta': json.dumps(delta),
            'option_text': ['الف', 'ب'], 'correct_options': ['1'],
        })
        self.assertEqual(response.status_code, 302)
        with connection.cursor() as cursor:
            cursor.execute("SELECT text, scoring_settings FROM questions WHERE type = 'single' AND teacher_id = %s", [self.teacher_in_id])
            text, settings_data = cursor.fetchone()
        self.assertEqual(text, 'کدام صحیح است؟')
        self.assertEqual(json.loads(settings_data)['text_delta'], delta)
        # Fill-in-the-blank now has the rich editor too; a legacy plain stem seeds it with its [n] markers.
        self.assertEqual(_question_form_data({'type': 'fill_blank', 'text': 'قلب [۱]', 'scoring_settings': {}})['text_delta'], {'ops': [{'insert': 'قلب [۱]\n'}]})

    def test_fill_blank_rich_stem_saves_and_renders_inputs_for_markers(self):
        from apps.core.views import _fill_blank_stem_html
        self.client.force_login(self.teacher_user)
        delta = {'ops': [{'insert': 'جمجمه از '}, {'insert': 'دو بخش', 'attributes': {'bold': True}}, {'insert': ' [۱] و [۲] است.\n'}]}
        response = self.client.post(f"{reverse('core:teacher_question_create')}?type=fill_blank", {
            'type': 'fill_blank', 'structure': 'independent', 'course_id': self.course_in_id, 'text': '', 'text_delta': json.dumps(delta),
            'blank_count': '2', 'blank_answers_1': 'نوروکرانیوم', 'blank_answers_2': 'ویسروکرانیوم',
        })
        self.assertEqual(response.status_code, 302)
        with connection.cursor() as cursor:
            cursor.execute("SELECT text, scoring_settings FROM questions WHERE type = 'fill_blank' AND teacher_id = %s", [self.teacher_in_id])
            text, settings_data = cursor.fetchone()
        self.assertEqual(text, 'جمجمه از دو بخش [۱] و [۲] است.')
        html = _fill_blank_stem_html(json.loads(settings_data), text, 'q1', ['نورو<x>', ''])
        self.assertIn('<strong>دو بخش</strong>', html)
        self.assertEqual(html.count('class="qx-blank"'), 2)
        self.assertIn('value="نورو&lt;x&gt;"', html)
        self.assertNotIn('[۱]', html)
        # A marker split by formatting no longer matches the plain text: fall back to the plain segments.
        split = {'text_delta': {'ops': [{'insert': 'قلب [۱'}, {'insert': ']', 'attributes': {'bold': True}}, {'insert': '\n'}]}}
        self.assertEqual(_fill_blank_stem_html(split, 'قلب [۱]', 'q1', []), '')

    def test_legacy_plain_essay_seeds_the_editor(self):
        from apps.core.views import _question_form_data
        form = _question_form_data({'type': 'essay', 'text': 'سؤال قدیمی', 'scoring_settings': {}})
        self.assertEqual(form['text_delta'], {'ops': [{'insert': 'سؤال قدیمی\n'}]})


class EssayAnswerFlowTests(_ScopedOrgFixture):
    """پیوست پاسخ تشریحی توسط دانشجو، دسترسی محدود به فایل‌ها و تصحیح واقعی با روبریک."""

    PNG = b'\x89PNG\r\n\x1a\n' + b'0' * 32

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        User = get_user_model()
        cls.teacher_user = User.objects.create_user(username='teacher_in_scope', email='t_in@example.com', password='x')
        cls.other_teacher_user = User.objects.create_user(username='teacher_out_scope', email='t_out@example.com', password='x')
        cls.student_user = User.objects.create_user(username='student_in_scope', email='student_in_scope@example.com', password='x')
        cls.other_student_user = User.objects.create_user(username='student_out_scope', email='student_out_scope@example.com', password='x')
        cls.question_id, cls.attempt_id = str(uuid.uuid4()), str(uuid.uuid4())
        _insert('questions', ['id', 'teacher_id', 'text', 'type', 'default_points', 'rubric', 'scoring_settings'], [
            cls.question_id, cls.teacher_in_id, 'سازوکار بارورفلکس را توضیح دهید', 'essay', 5,
            json.dumps([{'criterion': 'دقت علمی', 'points': '6'}, {'criterion': 'ساختار', 'points': '4'}]),
            json.dumps({'essay': {'answer_formats': ['text', 'image', 'file'], 'sample_answer_text': 'پاسخ مرجع'}}),
        ])
        _insert('exam_questions', ['id', 'exam_id', 'question_id', 'points', 'order_index'], [str(uuid.uuid4()), cls.exam_in_id, cls.question_id, 5, 1])
        _insert('exam_attempts', ['id', 'exam_id', 'student_id', 'started_at', 'status'], [cls.attempt_id, cls.exam_in_id, cls.student_in_id, timezone.now(), 'in_progress'])

    def setUp(self):
        self.private_root = tempfile.mkdtemp()
        override = self.settings(PRIVATE_MEDIA_ROOT=self.private_root)
        override.enable()
        self.addCleanup(override.disable)

    def _upload(self, **files):
        from django.core.files.uploadedfile import SimpleUploadedFile
        data = {f'answer_{self.question_id}': 'متن پاسخ'}
        for kind, (name, content) in files.items():
            data[f'upload_{self.question_id}_{kind}'] = SimpleUploadedFile(name, content)
        return self.client.post(reverse('core:student_attempt', args=[self.attempt_id]), data)

    def _files(self):
        with connection.cursor() as cursor:
            cursor.execute('SELECT answer_files FROM attempt_answers WHERE attempt_id = %s AND question_id = %s', [self.attempt_id, self.question_id])
            row = cursor.fetchone()
        return json.loads(row[0]) if row else []

    def _submit(self):
        with connection.cursor() as cursor:
            cursor.execute("UPDATE exam_attempts SET submitted_at = %s, status = 'submitted' WHERE id = %s", [timezone.now(), self.attempt_id])
            cursor.execute('UPDATE attempt_answers SET needs_manual_grading = %s WHERE attempt_id = %s', [True, self.attempt_id])

    def test_student_uploads_only_allowed_validated_files(self):
        self.client.force_login(self.student_user)
        self._upload(image=('scan.png', self.PNG), video=('clip.mp4', b'\x00\x00\x00\x18ftypmp42'))
        files = self._files()
        self.assertEqual([(item['kind'], item['name']) for item in files], [('image', 'scan.png')])  # video not allowed here
        from pathlib import Path
        self.assertTrue((Path(self.private_root) / files[0]['path']).is_file())
        response = self._upload(file=('notes.pdf', b'MZ not a pdf'))
        self.assertEqual(len(self._files()), 1)
        self.assertEqual(response.status_code, 302)
        page = self.client.get(response['Location'])
        self.assertContains(page, 'با پسوند آن هم‌خوانی ندارد')
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.post(reverse('core:student_attempt', args=[self.attempt_id]), {
            f'upload_{self.question_id}_image': SimpleUploadedFile('x.png', self.PNG),
        }, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(len(self._files()), 1)  # autosave never stores files
        self.client.post(reverse('core:student_attempt', args=[self.attempt_id]), {f'remove_file_{self.question_id}': files[0]['id']})
        self.assertEqual(self._files(), [])
        self.assertFalse((Path(self.private_root) / files[0]['path']).exists())

    def _answer(self, question_id):
        with connection.cursor() as cursor:
            cursor.execute('SELECT answer FROM attempt_answers WHERE attempt_id = %s AND question_id = %s', [self.attempt_id, question_id])
            row = cursor.fetchone()
        return json.loads(row[0])['value'] if row else None

    def test_answers_only_for_own_exam_questions_and_only_before_submission(self):
        url = reverse('core:student_attempt', args=[self.attempt_id])
        self.client.force_login(self.student_user)
        self.client.post(url, {f'answer_{self.question_id}': 'پاسخ اول', f'answer_{self.question_in_id}': 'نفوذ'})
        self.assertEqual(self._answer(self.question_id), 'پاسخ اول')
        self.assertIsNone(self._answer(self.question_in_id))  # question exists but is not part of this exam
        self.client.post(url, {f'answer_{self.question_id}': 'تغییر خودکار'}, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(self._answer(self.question_id), 'تغییر خودکار')
        self._submit()
        self.client.post(url, {f'answer_{self.question_id}': 'تغییر پس از ارسال'})
        self.client.post(url, {f'answer_{self.question_id}': 'تغییر پس از ارسال'}, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.client.post(reverse('core:student_attempt_submit', args=[self.attempt_id]), {f'answer_{self.question_id}': 'تغییر پس از ارسال'})
        self.assertEqual(self._answer(self.question_id), 'تغییر خودکار')

    def _started(self, minutes_ago, seconds_ago=0):
        from datetime import timedelta
        with connection.cursor() as cursor:
            cursor.execute('UPDATE exam_attempts SET started_at = %s WHERE id = %s',
                           [timezone.now() - timedelta(minutes=minutes_ago, seconds=seconds_ago), self.attempt_id])

    def test_answers_are_refused_after_the_time_limit(self):
        url = reverse('core:student_attempt', args=[self.attempt_id])
        self.client.force_login(self.student_user)
        self._started(90, 30)  # exam has no duration -> 90 minutes; still inside the 60s grace
        self.client.post(url, {f'answer_{self.question_id}': 'در مهلت'}, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(self._answer(self.question_id), 'در مهلت')
        self._started(92)
        late = self.client.post(url, {f'answer_{self.question_id}': 'دیرهنگام'}, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(late.status_code, 409)
        response = self.client.post(url, {f'answer_{self.question_id}': 'دیرهنگام'})
        self.assertRedirects(response, reverse('core:student_attempt_submit', args=[self.attempt_id]), fetch_redirect_response=False)
        self._upload(image=('late.png', self.PNG))
        self.client.post(reverse('core:student_attempt_submit', args=[self.attempt_id]), {f'answer_{self.question_id}': 'دیرهنگام'})
        self.assertEqual(self._answer(self.question_id), 'در مهلت')
        self.assertEqual(self._files(), [])

    def test_uploads_are_refused_after_submission(self):
        self.client.force_login(self.student_user)
        self._submit()
        self._upload(image=('late.png', self.PNG))
        self.assertEqual(self._files(), [])

    def test_answer_file_access_is_limited_to_student_and_grading_teacher(self):
        self.client.force_login(self.student_user)
        self._upload(image=('scan.png', self.PNG))
        url = reverse('core:attempt_answer_file', args=[self.attempt_id, self._files()[0]['id']])
        self.assertEqual(self.client.get(url).status_code, 200)
        self.client.force_login(self.teacher_user)
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['X-Content-Type-Options'], 'nosniff')
        for user in (self.other_teacher_user, self.other_student_user):
            self.client.force_login(user)
            self.assertEqual(self.client.get(url).status_code, 404)
        self.assertEqual(self.client.get(reverse('core:attempt_answer_file', args=[self.attempt_id, '..'])).status_code, 404)

    def test_teacher_grades_with_rubric_and_attempt_score_updates(self):
        self.client.force_login(self.student_user)
        self._upload(image=('scan.png', self.PNG))
        self._submit()
        with connection.cursor() as cursor:
            cursor.execute('SELECT id FROM attempt_answers WHERE attempt_id = %s', [self.attempt_id])
            answer_id = str(cursor.fetchone()[0])
        detail_url = reverse('core:teacher_grading_detail', args=[answer_id])

        self.client.force_login(self.other_teacher_user)
        self.assertNotContains(self.client.get(reverse('core:teacher_reviews')), 'سازوکار بارورفلکس')
        self.assertEqual(self.client.get(detail_url).status_code, 404)
        self.assertEqual(self.client.post(detail_url, {'score_0': '6', 'score_1': '4'}).status_code, 404)

        self.client.force_login(self.teacher_user)
        self.assertContains(self.client.get(reverse('core:teacher_reviews')), 'سازوکار بارورفلکس')
        page = self.client.get(detail_url)
        self.assertContains(page, 'متن پاسخ')
        self.assertContains(page, 'scan.png')
        self.assertContains(page, 'پاسخ مرجع')
        bad = self.client.post(detail_url, {'score_0': '9', 'score_1': '2'})
        self.assertEqual(bad.status_code, 200)
        self.assertContains(bad, 'باید بین ۰ و 6 باشد')
        self.client.post(detail_url, {'score_0': '6', 'score_1': '2', 'feedback': 'ساختار ضعیف بود'})
        with connection.cursor() as cursor:
            cursor.execute('SELECT points_awarded, needs_manual_grading, grader_feedback FROM attempt_answers WHERE id = %s', [answer_id])
            points, pending, feedback = cursor.fetchone()
            cursor.execute('SELECT score, status, is_graded FROM exam_attempts WHERE id = %s', [self.attempt_id])
            score, status, is_graded = cursor.fetchone()
        self.assertEqual(float(points), 4.0)  # 8 of 10 rubric points scaled to the question's 5 exam points
        self.assertFalse(pending)
        self.assertEqual(feedback, 'ساختار ضعیف بود')
        self.assertEqual((float(score), status, bool(is_graded)), (4.0, 'graded', True))


class QuestionTypesV2Tests(_ScopedOrgFixture):
    """سؤال‌های پاسخ عددی، ماتریسی، تصویری/نشانه‌گذاری و صوتی/ویدئویی: سازنده، تصحیح و صفحه‌ی آزمون دانشجو."""

    PNG = b'\x89PNG\r\n\x1a\n' + b'0' * 32
    MP3 = b'ID3' + b'0' * 32

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        User = get_user_model()
        cls.teacher_user = User.objects.create_user(username='teacher_in_scope', email='t_in@example.com', password='x')
        cls.student_user = User.objects.create_user(username='student_in_scope', email='student_in_scope@example.com', password='x')

    def setUp(self):
        for override in (self.settings(MEDIA_ROOT=tempfile.mkdtemp()), self.settings(PRIVATE_MEDIA_ROOT=tempfile.mkdtemp())):
            override.enable()
            self.addCleanup(override.disable)

    def _create(self, qtype, data):
        payload = {'type': qtype, 'structure': 'independent', 'course_id': self.course_in_id, 'is_published': '1', 'default_points': '2'}
        payload.update(data)
        return self.client.post(f"{reverse('core:teacher_question_create')}?type={qtype}", payload)

    def _row(self, text):
        with connection.cursor() as cursor:
            cursor.execute('SELECT id, type, options, correct_answer, accepted_answers, scoring_settings, rubric, question_media FROM questions WHERE text = %s', [text])
            row = cursor.fetchone()
        if not row:
            return None
        load = lambda value: json.loads(value) if isinstance(value, str) and value else value
        return {'id': row[0], 'type': row[1], 'options': load(row[2]), 'correct_answer': load(row[3]), 'accepted_answers': load(row[4]),
                'scoring_settings': load(row[5]), 'rubric': load(row[6]), 'question_media': load(row[7])}

    def _attempt(self, question_columns, answer_value=None):
        question_id, attempt_id = str(uuid.uuid4()), str(uuid.uuid4())
        columns = {'type': 'numeric', **question_columns}
        _insert('questions', ['id', 'teacher_id', 'course_id', 'text', 'default_points'] + list(columns),
                [question_id, self.teacher_in_id, self.course_in_id, 'سؤال نوع جدید', 2] + [
                    json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value for value in columns.values()])
        exam_id = str(uuid.uuid4())
        _insert('exams', ['id', 'title', 'teacher_id', 'course_id', 'start_at'], [exam_id, 'آزمون نوع جدید', self.teacher_in_id, self.course_in_id, timezone.now()])
        _insert('exam_questions', ['id', 'exam_id', 'question_id', 'points', 'order_index'], [str(uuid.uuid4()), exam_id, question_id, 2, 1])
        _insert('exam_attempts', ['id', 'exam_id', 'student_id', 'started_at', 'status'], [attempt_id, exam_id, self.student_in_id, timezone.now(), 'in_progress'])
        if answer_value is not None:
            _insert('attempt_answers', ['id', 'attempt_id', 'question_id', 'answer', 'needs_manual_grading'],
                    [str(uuid.uuid4()), attempt_id, question_id, json.dumps({'value': answer_value}, ensure_ascii=False), False])
        return question_id, attempt_id

    def _score(self, question_columns, answer_value):
        from apps.core.views import grade_attempt
        _question_id, attempt_id = self._attempt(question_columns, answer_value)
        return grade_attempt(attempt_id)

    # ----- builder -----

    def test_builder_pages_render_for_new_types_and_picker_links_them(self):
        from apps.core.views import _sx_type_label
        self.client.force_login(self.teacher_user)
        url = reverse('core:teacher_question_create')
        markers = {'numeric': 'name="numeric_answer"', 'matrix': 'data-matrix-table', 'image_label': 'data-label-canvas', 'media_response': 'name="media_max_plays"'}
        for kind, marker in markers.items():
            html = self.client.get(f'{url}?type={kind}').content.decode()
            self.assertIn(marker, html, kind)
            self.assertIn(f'سؤال ۱ — {_sx_type_label(kind)}', html, kind)
            self.assertIn('js/question-types-v2', html, kind)
        picker = self.client.get(reverse('core:teacher_question_type')).content.decode()
        for kind in markers:
            self.assertIn(f'?type={kind}"', picker, kind)
        self.assertIn('نقطه‌ای / مربوط به تصویر', picker)  # hotspot stays «در حال توسعه»

    def test_numeric_question_saves_tolerance_and_reloads_in_edit_form(self):
        self.client.force_login(self.teacher_user)
        response = self._create('numeric', {
            'text': 'حجم دارو را محاسبه کنید', 'numeric_answer': '۵', 'numeric_unit': 'mL', 'numeric_decimals': '1',
            'numeric_tolerance_mode': 'absolute', 'numeric_tolerance': '0.1', 'numeric_accept_persian': '1', 'numeric_require_unit': '1',
        })
        self.assertEqual(response.status_code, 302)
        row = self._row('حجم دارو را محاسبه کنید')
        self.assertEqual(row['type'], 'numeric')
        self.assertEqual(row['correct_answer'], '5')
        numeric = row['scoring_settings']['numeric']
        self.assertEqual((numeric['value'], numeric['unit'], numeric['tolerance'], numeric['require_unit']), (5.0, 'mL', 0.1, True))
        edit = self.client.get(reverse('core:teacher_question_edit', args=[row['id']])).content.decode()
        self.assertIn('name="numeric_answer" value="5"', edit)
        self.assertIn('name="numeric_unit" value="mL"', edit)

    def test_numeric_question_validation(self):
        self.client.force_login(self.teacher_user)
        response = self._create('numeric', {'text': 'عدد نامعتبر', 'numeric_answer': 'پنج', 'numeric_require_unit': '1'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'پاسخ صحیح عددی را به‌صورت عدد وارد کنید.')
        self.assertContains(response, 'واحد پاسخ را وارد کنید')
        self.assertIsNone(self._row('عدد نامعتبر'))

    def test_matrix_question_saves_rows_with_server_ids(self):
        self.client.force_login(self.teacher_user)
        data = {
            'columns': [{'id': 'cx', 'text': 'پرکاری تیروئید'}, {'id': 'cy', 'text': 'کم‌کاری تیروئید'}],
            'rows': [
                {'id': 'ra', 'text': 'تاکی‌کاردی', 'points': 1, 'correct': ['cx']},
                {'id': 'rb', 'text': 'افزایش وزن', 'points': 2, 'correct': ['cy', 'bogus']},
            ],
        }
        response = self._create('matrix', {'text': 'علائم تیروئید', 'matrix_json': json.dumps(data, ensure_ascii=False),
                                           'matrix_selection': 'single', 'matrix_show_column_titles': '1', 'scoring_mode': 'partial'})
        self.assertEqual(response.status_code, 302)
        row = self._row('علائم تیروئید')
        self.assertEqual([(item['id'], item['correct'], item['points']) for item in row['options']], [('r1', ['c1'], 1.0), ('r2', ['c2'], 2.0)])
        self.assertEqual(row['correct_answer'], ['r1:c1', 'r2:c2'])
        self.assertEqual(row['scoring_settings']['matrix']['columns'][0], {'id': 'c1', 'text': 'پرکاری تیروئید'})

    def test_matrix_validation(self):
        self.client.force_login(self.teacher_user)
        data = {'columns': [{'id': 'a', 'text': 'الف'}], 'rows': [{'id': 'r', 'text': 'سطر', 'correct': []}]}
        response = self._create('matrix', {'text': 'ماتریس ناقص', 'matrix_json': json.dumps(data), 'matrix_selection': 'single'})
        self.assertContains(response, 'حداقل دو ستون با عنوان وارد کنید.')
        self.assertContains(response, 'حداقل دو سطر با متن وارد کنید.')
        self.assertContains(response, 'برای هر سطر پاسخ صحیح را در جدول مشخص کنید.')
        two = {'columns': [{'id': 'a', 'text': 'الف'}, {'id': 'b', 'text': 'ب'}],
               'rows': [{'id': 'r1', 'text': 'یک', 'correct': ['a', 'b']}, {'id': 'r2', 'text': 'دو', 'correct': ['a']}]}
        response = self._create('matrix', {'text': 'ماتریس تک‌گزینه‌ای', 'matrix_json': json.dumps(two), 'matrix_selection': 'single'})
        self.assertContains(response, 'در حالت تک‌گزینه‌ای هر سطر فقط یک پاسخ صحیح دارد.')
        self.assertIsNone(self._row('ماتریس تک‌گزینه‌ای'))

    def test_image_label_requires_image_and_saves_markers(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.force_login(self.teacher_user)
        markers = [{'x': 40.5, 'y': 20, 'answer': 'آئورت', 'synonyms': 'آئورت صعودی، سرخرگ آئورت', 'points': 1},
                   {'x': 150, 'y': 70, 'answer': 'بطن چپ', 'synonyms': '', 'points': 2}]
        response = self._create('image_label', {'text': 'بخش‌های قلب', 'label_markers_json': json.dumps(markers, ensure_ascii=False)})
        self.assertContains(response, 'تصویر سؤال را برای نشانه‌گذاری بارگذاری کنید.')
        response = self._create('image_label', {
            'text': 'بخش‌های قلب', 'label_markers_json': json.dumps(markers, ensure_ascii=False), 'label_answer_mode': 'drag',
            'label_accept_synonyms': '1', 'question_image_file': SimpleUploadedFile('heart.png', self.PNG, content_type='image/png'),
        })
        self.assertEqual(response.status_code, 302)
        row = self._row('بخش‌های قلب')
        self.assertTrue(row['question_media']['image'].endswith('.png'))
        self.assertEqual([(m['marker'], m['x'], m['answer'], m['synonyms']) for m in row['accepted_answers']],
                         [(1, 40.5, 'آئورت', ['آئورت صعودی', 'سرخرگ آئورت']), (2, 100.0, 'بطن چپ', [])])
        self.assertEqual(row['scoring_settings']['image_label'], {'answer_mode': 'drag', 'accept_synonyms': True})

    def test_media_response_requires_media_and_saves_playback_rubric_and_attachment(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.force_login(self.teacher_user)
        base = {
            'text': 'صدای قلب را توضیح دهید', 'media_allow_replay': '1', 'media_range_start': '00:10', 'media_range_end': '00:40',
            'media_max_plays': '3', 'media_transcript': 'صدای قلب بیمار ۶۵ ساله', 'media_show_transcript': '1',
            'answer_formats': ['audio', 'text', 'image'], 'media_max_record_seconds': '120',
            'media_rubric_criterion': ['تشخیص صدا', 'استدلال بالینی'], 'media_rubric_description': ['شناسایی سوفل', 'ارتباط با بیمار'],
            'media_rubric_points': ['2', '3'],
        }
        response = self._create('media_response', base)
        self.assertContains(response, 'فایل صوتی یا ویدئویی سؤال را بارگذاری کنید.')
        response = self._create('media_response', {**base, 'media_range_end': '00:05'})
        self.assertContains(response, 'پایان بازه‌ی مجاز پخش باید بعد از شروع آن باشد.')
        bad = SimpleUploadedFile('notes.exe', b'MZ')
        response = self._create('media_response', {**base, 'question_audio_file': SimpleUploadedFile('heart.mp3', self.MP3), 'media_attachment_file': bad})
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(self._row('صدای قلب را توضیح دهید'))
        response = self._create('media_response', {
            **base, 'question_audio_file': SimpleUploadedFile('heart.mp3', self.MP3),
            'media_attachment_file': SimpleUploadedFile('guide.pdf', b'%PDF-1.4 test'),
        })
        self.assertEqual(response.status_code, 302)
        row = self._row('صدای قلب را توضیح دهید')
        media = row['scoring_settings']['media_response']
        self.assertEqual((media['range_start'], media['range_end'], media['max_plays'], media['max_record_seconds']), (10, 40, 3, 120))
        self.assertEqual(media['attachment']['name'], 'guide.pdf')
        self.assertEqual(row['scoring_settings']['essay']['answer_formats'], ['text', 'audio'])  # image is not offered here
        self.assertEqual([(item['criterion'], item['points']) for item in row['rubric']], [('تشخیص صدا', '2'), ('استدلال بالینی', '3')])
        edit = self.client.get(reverse('core:teacher_question_edit', args=[row['id']])).content.decode()
        self.assertIn('name="media_range_start" value="00:10"', edit)
        self.assertIn('guide.pdf', edit)

    # ----- grading -----

    def test_numeric_grading_tolerance_digits_and_unit(self):
        numeric = {'value': 5, 'unit': 'mL', 'tolerance_mode': 'absolute', 'tolerance': 0.1, 'accept_persian_digits': True, 'require_unit': False}
        settings_of = lambda **changes: {'scoring_settings': {'numeric': {**numeric, **changes}}}
        self.assertEqual(self._score(settings_of(), '5.08')['score'], 2)
        self.assertEqual(self._score(settings_of(), '۵٫۰۵')['score'], 2)
        self.assertEqual(self._score(settings_of(), '5.2')['score'], 0)
        self.assertEqual(self._score(settings_of(accept_persian_digits=False), '۵')['score'], 0)
        self.assertEqual(self._score(settings_of(tolerance_mode='percent', tolerance=10), '5.4')['score'], 2)
        self.assertEqual(self._score(settings_of(require_unit=True), ['5', 'ml'])['score'], 2)
        self.assertEqual(self._score(settings_of(require_unit=True), ['5', 'L'])['score'], 0)
        self.assertEqual(self._score(settings_of(require_unit=True), '5')['score'], 0)

    def test_matrix_grading_full_partial_and_all_or_nothing(self):
        rows = [{'id': 'r1', 'text': 'الف', 'points': 1, 'correct': ['c1']}, {'id': 'r2', 'text': 'ب', 'points': 3, 'correct': ['c2']}]
        question = lambda mode: {'type': 'matrix', 'options': rows, 'scoring_settings': {'scoring_mode': mode, 'matrix': {'columns': [{'id': 'c1'}, {'id': 'c2'}]}}}
        self.assertEqual(self._score(question('partial'), 'r1:c1,r2:c2')['score'], 2)
        self.assertEqual(self._score(question('partial'), 'r1:c2,r2:c2')['score'], 1.5)  # 3 of 4 row weight
        self.assertEqual(self._score(question('all_or_nothing'), 'r1:c2,r2:c2')['score'], 0)
        self.assertEqual(self._score(question('partial'), 'r1:c1,r1:c2,r2:c2')['score'], 1.5)  # extra pick breaks row 1

    def test_image_label_grading_synonyms_and_partial(self):
        markers = [{'marker': 1, 'x': 10, 'y': 10, 'answer': 'آئورت', 'synonyms': ['سرخرگ آئورت'], 'points': 1},
                   {'marker': 2, 'x': 50, 'y': 50, 'answer': 'بطن چپ', 'synonyms': [], 'points': 1}]
        question = lambda synonyms, mode='partial': {'type': 'image_label', 'accepted_answers': markers,
                                                     'scoring_settings': {'scoring_mode': mode, 'image_label': {'accept_synonyms': synonyms}}}
        self.assertEqual(self._score(question(True), ['سرخرگ آئورت', 'بطن  چپ'])['score'], 2)
        self.assertEqual(self._score(question(False), ['سرخرگ آئورت', 'بطن چپ'])['score'], 1)
        self.assertEqual(self._score(question(False, 'all_or_nothing'), ['سرخرگ آئورت', 'بطن چپ'])['score'], 0)

    def test_media_response_goes_to_manual_grading_queue(self):
        result = self._score({'type': 'media_response', 'scoring_settings': {'essay': {'answer_formats': ['text']}}}, 'پاسخ متنی')
        self.assertEqual(result['score'], 0)
        with connection.cursor() as cursor:
            cursor.execute("SELECT aa.needs_manual_grading FROM attempt_answers aa JOIN questions q ON q.id = aa.question_id WHERE q.type = 'media_response'")
            self.assertTrue(cursor.fetchone()[0])
        self.client.force_login(self.teacher_user)
        self.assertContains(self.client.get(reverse('core:teacher_reviews')), 'سؤال نوع جدید')

    # ----- student exam page -----

    def test_student_page_renders_matrix_without_correct_answers(self):
        rows = [{'id': 'r1', 'text': 'تاکی‌کاردی', 'points': 1, 'correct': ['c1']}, {'id': 'r2', 'text': 'افزایش وزن', 'points': 1, 'correct': ['c2']}]
        _question_id, attempt_id = self._attempt({'type': 'matrix', 'options': rows, 'scoring_settings': {'matrix': {
            'columns': [{'id': 'c1', 'text': 'پرکاری'}, {'id': 'c2', 'text': 'کم‌کاری'}], 'selection': 'single',
            'show_column_titles': True, 'show_row_titles': True}}}, 'r1:c1')
        self.client.force_login(self.student_user)
        page = self.client.get(reverse('core:student_attempt', args=[attempt_id])).content.decode()
        self.assertIn('data-matrix', page)
        self.assertIn('<th scope="col">کم‌کاری</th>', page)
        self.assertIn('value="r1:c1" data-matrix-cell checked', page)
        self.assertIn('value="r1:c1" data-matrix-answer', page)
        self.assertNotIn('"correct"', page)

    def test_student_page_saves_numeric_answer_with_unit(self):
        question_id, attempt_id = self._attempt({'scoring_settings': {'numeric': {'value': 5, 'unit': 'mL', 'require_unit': True}}})
        self.client.force_login(self.student_user)
        url = reverse('core:student_attempt', args=[attempt_id])
        self.assertContains(self.client.get(url), 'aria-label="واحد پاسخ"')
        self.client.post(url, {f'answer_{question_id}': ['۵', 'mL']})
        page = self.client.get(url).content.decode()
        self.assertIn('value="۵"', page)
        self.assertIn('value="mL"', page)

    def test_student_page_renders_image_label_pins_and_drag_bank(self):
        markers = [{'marker': 1, 'x': 33.3333, 'y': 20, 'answer': 'آئورت', 'synonyms': ['سرخرگ آئورت'], 'points': 1}]
        _question_id, attempt_id = self._attempt({'type': 'image_label', 'accepted_answers': markers,
                                                  'question_media': {'image': '/media/question-media/heart.png'},
                                                  'scoring_settings': {'image_label': {'answer_mode': 'drag'}}})
        self.client.force_login(self.student_user)
        page = self.client.get(reverse('core:student_attempt', args=[attempt_id])).content.decode()
        self.assertIn('style="left:33.33%;top:20.00%"', page)
        self.assertIn('data-label-chip>آئورت</button>', page)
        self.assertNotIn('سرخرگ آئورت', page)  # synonyms never reach the student
        self.assertEqual(page.count('/media/question-media/heart.png'), 1)  # not repeated as a plain figure

    def test_student_media_response_page_and_recorded_audio_upload(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        question_id, attempt_id = self._attempt({'type': 'media_response', 'question_media': {'audio': '/media/question-media/heart.mp3'},
                                                 'scoring_settings': {'essay': {'answer_formats': ['audio', 'video']}, 'media_response': {
                                                     'allow_replay': True, 'max_plays': 3, 'range_start': 10, 'range_end': 40,
                                                     'transcript': 'متن گفتار', 'show_transcript': True, 'max_record_seconds': 90}}})
        self.client.force_login(self.student_user)
        url = reverse('core:student_attempt', args=[attempt_id])
        page = self.client.get(url).content.decode()
        self.assertIn('data-playback', page)
        self.assertIn('"max_plays": 3', page)
        self.assertIn('data-record-kind="video" data-record-max="90"', page)
        self.assertIn('متن گفتار', page)
        self.client.post(url, {f'answer_{question_id}': '', f'upload_{question_id}_audio': SimpleUploadedFile('recording.mp3', self.MP3)})
        with connection.cursor() as cursor:
            cursor.execute('SELECT answer_files FROM attempt_answers WHERE attempt_id = %s AND question_id = %s', [attempt_id, question_id])
            files = json.loads(cursor.fetchone()[0])
        self.assertEqual([(item['kind'], item['name']) for item in files], [('audio', 'recording.mp3')])


class TeacherDashboardTests(_ScopedOrgFixture):
    """داشبورد استاد فقط داده‌ی آزمون‌ها و سؤال‌های همان استاد را نشان می‌دهد."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        User = get_user_model()
        cls.teacher_user = User.objects.create_user(username='teacher_in_scope', email='t_in@example.com', password='x')

    def _pending_answer(self, exam_id, question_id):
        attempt_id = str(uuid.uuid4())
        _insert('exam_attempts', ['id', 'exam_id', 'student_id', 'started_at', 'status'],
                [attempt_id, exam_id, self.student_in_id, timezone.now(), 'submitted'])
        _insert('attempt_answers', ['id', 'attempt_id', 'question_id', 'answer', 'needs_manual_grading'],
                [str(uuid.uuid4()), attempt_id, question_id, json.dumps({'value': 'x'}), True])

    def test_pending_review_count_excludes_other_teachers_answers(self):
        self._pending_answer(self.exam_out_id, self.question_out_id)
        self._pending_answer(self.exam_out_id, self.question_out_id)
        self._pending_answer(self.exam_in_id, self.question_in_id)
        self.client.force_login(self.teacher_user)
        response = self.client.get(reverse('core:teacher_panel'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['teacher_stats']['reviews'], 1)

    def _exam(self, teacher_id, title, start_at, **columns):
        exam_id = str(uuid.uuid4())
        _insert('exams', ['id', 'title', 'teacher_id', 'course_id', 'start_at', 'duration_minutes'] + list(columns),
                [exam_id, title, teacher_id, self.course_in_id, start_at, 60] + list(columns.values()))
        return exam_id

    def test_dashboard_lists_only_own_upcoming_exams_with_participants(self):
        from datetime import timedelta
        soon = timezone.now() + timedelta(days=3)
        own_id = self._exam(self.teacher_in_id, 'آزمون آینده من', soon, lifecycle_status='published')
        self._exam(self.teacher_out_id, 'آزمون آینده دیگری', soon, lifecycle_status='published')
        _insert('student_group_members', ['id', 'group_id', 'student_user_id', 'full_name', 'national_id', 'student_number'],
                [str(uuid.uuid4()), self.group_in_id, self.student_in_id, 'دانشجو', '1', '1'])
        _insert('exam_assignments', ['id', 'exam_id', 'group_id'], [str(uuid.uuid4()), own_id, self.group_in_id])
        self.client.force_login(self.teacher_user)
        response = self.client.get(reverse('core:teacher_panel'))
        data = response.context['dashboard']
        self.assertEqual([row['title'] for row in data['upcoming_exams']], ['آزمون آینده من'])
        self.assertEqual(data['upcoming_count'], 1)
        self.assertEqual(data['upcoming_exams'][0]['participants'], 1)
        self.assertEqual(data['upcoming_exams'][0]['status_label'], 'برنامه‌ریزی شده')
        self.assertEqual(data['next_exam']['title'], 'آزمون آینده من')
        page = response.content.decode()
        self.assertIn('آزمون آینده من', page)
        self.assertNotIn('آزمون آینده دیگری', page)
        for section in ('دسترسی سریع', 'آزمون‌های پیش رو', 'یادآور آزمون آینده', 'فعالیت‌های اخیر', 'عملکرد دانشجویان', 'پاسخ‌های نیازمند تصحیح'):
            self.assertIn(section, page)
        self.assertNotIn('۷۸٪', page)  # the old hard-coded participation figure is gone

    def test_dashboard_pending_and_performance_are_scoped(self):
        self._pending_answer(self.exam_in_id, self.question_in_id)
        self._pending_answer(self.exam_in_id, self.question_in2_id)
        self._pending_answer(self.exam_out_id, self.question_out_id)
        for exam_id, score in ((self.exam_in_id, 15), (self.exam_in_id, 17), (self.exam_out_id, 2)):
            _insert('exam_attempts', ['id', 'exam_id', 'student_id', 'started_at', 'status', 'score', 'max_score'],
                    [str(uuid.uuid4()), exam_id, self.student_in_id, timezone.now(), 'graded', score, 20])
        self.client.force_login(self.teacher_user)
        data = self.client.get(reverse('core:teacher_panel')).context['dashboard']
        self.assertEqual(data['pending_total'], 2)
        self.assertEqual(data['pending_exam_count'], 1)
        self.assertEqual([(row['title'], row['count']) for row in data['pending_by_exam']], [('آزمون داخل محدوده', 2)])
        self.assertEqual([(row['title'], row['average']) for row in data['performance']], [('آزمون داخل محدوده', 16.0)])

    def test_jalali_long_date(self):
        from datetime import date
        from apps.core.views import _jalali_long_date
        self.assertEqual(_jalali_long_date(date(2025, 5, 19)), 'دوشنبه ۲۹ اردیبهشت ۱۴۰۴')
        self.assertEqual(_jalali_long_date(None), '')


class TeacherExamWizardTests(_ScopedOrgFixture):
    """
    استاد همان ویزارد ایجاد آزمون مدیر را می‌بیند ولی فقط با گروه‌ها، دانشجویان و
    سؤال‌های خودش؛ مالک آزمون همیشه خود اوست، نمی‌تواند منتشر کند، و همه‌ی
    انتخاب‌ها هنگام ثبت دوباره سمت سرور بررسی می‌شوند.
    """

    NID_IN = 'NID-TW-IN'
    NID_OUT = 'NID-TW-OUT'

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        User = get_user_model()
        cls.teacher_user = User.objects.create_user(username='teacher_in_scope', email='t_in@example.com', password='x')
        for student_id, group_id, nid in ((cls.student_in_id, cls.group_in_id, cls.NID_IN), (cls.student_out_id, cls.group_out_id, cls.NID_OUT)):
            _insert('student_group_members', ['id', 'group_id', 'student_user_id', 'full_name', 'national_id', 'student_number'],
                    [str(uuid.uuid4()), group_id, student_id, 'دانشجو', nid, nid])
            with connection.cursor() as cursor:
                cursor.execute('UPDATE profiles SET national_id = %s WHERE id = %s', [nid, student_id])

    def _payload(self, **overrides):
        payload = {
            'title': 'آزمون استاد',
            'course_id': self.course_in_id,
            'group_ids': [self.group_in_id],
            'question_ids': [self.question_in_id, self.question_in2_id],
            'duration_minutes': '60',
            'start_at': '2026-11-10T09:00',
            'publish_choice': 'draft',
        }
        payload.update(overrides)
        return payload

    def _exam(self, title='آزمون استاد'):
        with connection.cursor() as cursor:
            cursor.execute('SELECT id, teacher_id, lifecycle_status, approval_status, is_published FROM exams WHERE title = %s', [title])
            return cursor.fetchall()

    def test_get_offers_only_own_resources(self):
        self.client.force_login(self.teacher_user)
        response = self.client.get(reverse('core:teacher_exam_create'))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'super_admin/exam_wizard.html')
        ctx = response.context
        self.assertEqual(ctx['wizard_role'], 'teacher')
        self.assertEqual({g['id'] for g in ctx['groups']}, {self.group_in_id})
        self.assertEqual({c['id'] for c in ctx['courses']}, {self.course_in_id})
        self.assertEqual({q['id'] for q in ctx['questions']}, {self.question_in_id, self.question_in2_id})
        self.assertEqual({s['id'] for s in ctx['students']}, {self.student_in_id})
        self.assertEqual({t['id'] for t in ctx['teachers']}, {self.teacher_in_id})
        page = response.content.decode()
        self.assertIn('ارسال برای تأیید', page)
        self.assertNotIn('value="immediate"><strong>', page)
        self.assertIn(reverse('core:teacher_exam_participants_match'), page)

    def test_draft_is_owned_by_teacher_even_if_another_teacher_is_posted(self):
        self.client.force_login(self.teacher_user)
        response = self.client.post(reverse('core:teacher_exam_create'), self._payload(teacher_id=self.teacher_out_id))
        self.assertRedirects(response, reverse('core:teacher_exams'), fetch_redirect_response=False)
        [(exam_id, teacher_id, lifecycle, approval, is_published)] = self._exam()
        self.assertEqual((teacher_id, lifecycle, approval, bool(is_published)), (self.teacher_in_id, 'draft', 'pending', False))
        with connection.cursor() as cursor:
            cursor.execute('SELECT question_id FROM exam_questions WHERE exam_id = %s ORDER BY order_index', [exam_id])
            self.assertEqual([row[0] for row in cursor.fetchall()], [self.question_in_id, self.question_in2_id])
            cursor.execute('SELECT group_id FROM exam_assignments WHERE exam_id = %s', [exam_id])
            self.assertEqual([row[0] for row in cursor.fetchall()], [self.group_in_id])

    def test_submit_sends_for_approval_and_cannot_publish(self):
        self.client.force_login(self.teacher_user)
        self.client.post(reverse('core:teacher_exam_create'), self._payload(publish_choice='submit'))
        [(_id, _teacher, lifecycle, approval, is_published)] = self._exam()
        self.assertEqual((lifecycle, approval, bool(is_published)), ('pending_approval', 'pending', False))
        self.client.post(reverse('core:teacher_exam_create'), self._payload(title='آزمون فوری', publish_choice='immediate'))
        [(_id, _teacher, lifecycle, approval, is_published)] = self._exam('آزمون فوری')
        self.assertEqual((lifecycle, approval, bool(is_published)), ('draft', 'pending', False))

    def test_out_of_scope_selections_are_rejected(self):
        self.client.force_login(self.teacher_user)
        for overrides in (
            {'course_id': self.course_out_id},
            {'group_ids': [self.group_in_id, self.group_out_id]},
            {'student_ids': [self.student_out_id]},
            {'excluded_student_ids': [self.student_out_id]},
            {'question_ids': [self.question_in_id, self.question_out_id]},
            {'staff_teacher_id': [self.teacher_out_id], 'staff_role': ['observer']},
        ):
            response = self.client.post(reverse('core:teacher_exam_create'), self._payload(**overrides))
            self.assertEqual(response.status_code, 403, overrides)
        self.assertEqual(self._exam(), [])

    def test_missing_audience_or_questions_shows_error(self):
        self.client.force_login(self.teacher_user)
        response = self.client.post(reverse('core:teacher_exam_create'), self._payload(group_ids=[]))
        self.assertRedirects(response, reverse('core:teacher_exam_create'), fetch_redirect_response=False)
        response = self.client.post(reverse('core:teacher_exam_create'), self._payload(question_ids=[]))
        self.assertRedirects(response, reverse('core:teacher_exam_create'), fetch_redirect_response=False)
        self.assertEqual(self._exam(), [])

    def test_participants_match_is_limited_to_own_group_members(self):
        self.client.force_login(self.teacher_user)
        upload = io.BytesIO(f'کد ملی یا شماره دانشجویی\n{self.NID_IN}\n{self.NID_OUT}\n'.encode('utf-8-sig'))
        upload.name = 'participants.csv'
        data = self.client.post(reverse('core:teacher_exam_participants_match'), {'excel_file': upload}).json()
        self.assertEqual([row['id'] for row in data['found']], [self.student_in_id])
        self.assertEqual(data['not_found'], [self.NID_OUT])

    def test_wizard_is_teacher_only(self):
        self.client.force_login(self.manager_user)
        self.assertEqual(self.client.get(reverse('core:teacher_exam_create')).status_code, 403)
        self.assertEqual(self.client.post(reverse('core:teacher_exam_participants_match')).status_code, 403)


class TeacherExamListTests(_ScopedOrgFixture):
    """صفحه‌ی «آزمون‌های من» فقط آزمون‌های خود استاد را با شمارش‌های واقعی، جستجو، فیلتر و صفحه‌بندی نشان می‌دهد."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        User = get_user_model()
        cls.teacher_user = User.objects.create_user(username='teacher_in_scope', email='t_in@example.com', password='x')

    def _get(self, **params):
        self.client.force_login(self.teacher_user)
        response = self.client.get(reverse('core:teacher_exams'), params)
        self.assertEqual(response.status_code, 200)
        return response

    def _exam(self, title, **columns):
        exam_id = str(uuid.uuid4())
        _insert('exams', ['id', 'title', 'teacher_id', 'course_id'] + list(columns),
                [exam_id, title, self.teacher_in_id, self.course_in_id] + list(columns.values()))
        return exam_id

    def test_lists_only_own_exams_with_real_counts(self):
        _insert('exam_questions', ['id', 'exam_id', 'question_id', 'points', 'order_index'], [str(uuid.uuid4()), self.exam_in_id, self.question_in_id, 1, 1])
        _insert('exam_questions', ['id', 'exam_id', 'question_id', 'points', 'order_index'], [str(uuid.uuid4()), self.exam_in_id, self.question_in2_id, 1, 2])
        _insert('student_group_members', ['id', 'group_id', 'student_user_id', 'full_name', 'national_id', 'student_number'],
                [str(uuid.uuid4()), self.group_in_id, self.student_in_id, 'دانشجو', '1', '1'])
        _insert('exam_assignments', ['id', 'exam_id', 'group_id'], [str(uuid.uuid4()), self.exam_in_id, self.group_in_id])
        attempt_id = str(uuid.uuid4())
        _insert('exam_attempts', ['id', 'exam_id', 'student_id', 'started_at', 'status'], [attempt_id, self.exam_in_id, self.student_in_id, timezone.now(), 'submitted'])
        _insert('attempt_answers', ['id', 'attempt_id', 'question_id', 'answer', 'needs_manual_grading'], [str(uuid.uuid4()), attempt_id, self.question_in_id, '{}', True])
        response = self._get()
        rows = response.context['exam_list']['rows']
        self.assertEqual([row['title'] for row in rows], ['آزمون داخل محدوده'])
        row = rows[0]
        self.assertEqual((row['question_count'], row['assigned'], row['submitted'], row['pending_reviews']), (2, 1, 1, 1))
        page = response.content.decode()
        self.assertNotIn('آزمون خارج محدوده', page)
        for leftover in ('ریاضی دهم', 'گروه الف', '۲۸/۳۰'):
            self.assertNotIn(leftover, page)
        self.assertIn('تصحیح (۱)', page)

    def test_status_filter_and_search(self):
        self._exam('پیش‌نویس فیزیولوژی', lifecycle_status='draft')
        self._exam('ارسالی فیزیولوژی', lifecycle_status='pending_approval')
        self._exam('لغوشده قلب', lifecycle_status='published', is_cancelled=True)
        titles = lambda response: {row['title'] for row in response.context['exam_list']['rows']}
        self.assertEqual(titles(self._get(status='pending_approval')), {'ارسالی فیزیولوژی'})
        self.assertEqual(titles(self._get(status='cancelled')), {'لغوشده قلب'})
        self.assertEqual(titles(self._get(q='فیزیولوژی')), {'پیش‌نویس فیزیولوژی', 'ارسالی فیزیولوژی'})
        self.assertEqual(titles(self._get(status='bogus', q='قلب')), {'لغوشده قلب'})  # unknown filter falls back to all
        self.assertIn('آزمونی با این جستجو یا وضعیت پیدا نشد', self._get(q='ناموجود').content.decode())

    def test_pagination(self):
        for index in range(24):
            self._exam(f'آزمون شماره {index}')
        first, second = self._get(), self._get(page='2')
        self.assertEqual(first.context['exam_list']['total'], 25)
        self.assertEqual(len(first.context['exam_list']['rows']), 20)
        self.assertEqual(len(second.context['exam_list']['rows']), 5)
        self.assertIn('صفحه ۱ از ۲', first.content.decode())

    def test_status_labels_follow_time_window(self):
        from datetime import timedelta
        from apps.core.views import _teacher_exam_status
        now = timezone.now()
        hour = timedelta(hours=1)
        self.assertEqual(_teacher_exam_status({'lifecycle': 'published', 'start_at': now - hour, 'end_at': now + hour}, now)[0], 'در حال برگزاری')
        self.assertEqual(_teacher_exam_status({'lifecycle': 'published', 'start_at': now + hour}, now)[0], 'برنامه‌ریزی شده')
        self.assertEqual(_teacher_exam_status({'lifecycle': 'published', 'start_at': now - 2 * hour, 'end_at': now - hour}, now)[0], 'پایان‌یافته')
        self.assertEqual(_teacher_exam_status({'lifecycle': 'pending_approval'}, now)[0], 'در انتظار تأیید')
        self.assertEqual(_teacher_exam_status({'lifecycle': 'published', 'is_cancelled': True}, now)[0], 'لغو شده')
        self.assertEqual(_teacher_exam_status({'lifecycle': ''}, now)[0], 'پیش‌نویس')


class TeacherClassListTests(_ScopedOrgFixture):
    """صفحه‌ی «کلاس‌های من» فقط گروه‌های خود استاد (مالک یا هم‌تدریس) را با آمار واقعی، جستجو، فیلتر و صفحه‌بندی نشان می‌دهد."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        User = get_user_model()
        cls.teacher_user = User.objects.create_user(username='teacher_in_scope', email='t_in@example.com', password='x')

    def _get(self, **params):
        self.client.force_login(self.teacher_user)
        response = self.client.get(reverse('core:teacher_courses'), params)
        self.assertEqual(response.status_code, 200)
        return response

    def _group(self, name, teacher_id=None, **columns):
        group_id = str(uuid.uuid4())
        columns = {'academic_year': '1404', **columns}
        _insert('student_groups', ['id', 'teacher_id', 'course_id', 'course_name'] + list(columns),
                [group_id, teacher_id or self.teacher_in_id, self.course_in_id, name] + list(columns.values()))
        return group_id

    def _member(self, group_id, student_id, name):
        _insert('student_group_members', ['id', 'group_id', 'student_user_id', 'full_name', 'national_id', 'student_number'],
                [str(uuid.uuid4()), group_id, student_id, name, '1', '1'])

    def test_lists_own_and_co_taught_groups_with_real_figures(self):
        shared_id = self._group('کلاس هم‌تدریس', teacher_id=self.teacher_out_id)
        _insert('group_teachers', ['group_id', 'teacher_id'], [shared_id, self.teacher_in_id])
        self._member(self.group_in_id, self.student_in_id, 'دانشجوی الف')
        self._member(shared_id, self.student_in_id, 'دانشجوی الف')
        self._member(shared_id, self.student_out_id, 'دانشجوی ب')
        _insert('exam_assignments', ['id', 'exam_id', 'group_id'], [str(uuid.uuid4()), self.exam_in_id, self.group_in_id])
        response = self._get()
        rows = {row['course_name']: row for row in response.context['class_list']['rows']}
        self.assertEqual(set(rows), {'درس داخل محدوده', 'کلاس هم‌تدریس'})
        self.assertEqual((rows['درس داخل محدوده']['students_count'], rows['درس داخل محدوده']['exams_count']), (1, 1))
        self.assertEqual(rows['کلاس هم‌تدریس']['students_count'], 2)
        self.assertEqual(response.context['class_list']['summary']['students'], 2)  # distinct across classes
        page = response.content.decode()
        self.assertNotIn('درس خارج محدوده', page)
        for leftover in ('مبانی پرستاری', 'فارماکولوژی', '۱۴۰۴/۰۳/۲۰'):
            self.assertNotIn(leftover, page)

    def test_search_term_and_status_filters(self):
        self._group('فیزیولوژی پایه', semester='اول', status='active')
        self._group('فیزیولوژی پیشرفته', semester='دوم', status='draft')
        self._group('آناتومی', academic_year='1403', semester='دوم', is_active=False)
        names = lambda response: {row['course_name'] for row in response.context['class_list']['rows']}
        self.assertEqual(names(self._get(q='فیزیولوژی')), {'فیزیولوژی پایه', 'فیزیولوژی پیشرفته'})
        self.assertEqual(names(self._get(term='1404|دوم')), {'فیزیولوژی پیشرفته'})
        # the fixture group keeps the column default status 'draft', as on the admin groups page
        self.assertEqual(names(self._get(status='inactive')), {'فیزیولوژی پیشرفته', 'آناتومی', 'درس داخل محدوده'})
        self.assertNotIn('آناتومی', names(self._get(status='active')))
        self.assertIn('فیزیولوژی پایه', names(self._get(term='9999|bogus', status='bogus')))  # unknown filters fall back to all
        self.assertIn('کلاسی با این جستجو یا فیلتر پیدا نشد', self._get(q='ناموجود').content.decode())

    def test_pagination_and_list_view(self):
        for index in range(13):
            self._group(f'کلاس شماره {index}')
        first, second = self._get(), self._get(page='2')
        self.assertEqual(first.context['class_list']['total'], 14)
        self.assertEqual(len(first.context['class_list']['rows']), 12)
        self.assertEqual(len(second.context['class_list']['rows']), 2)
        self.assertIn('صفحه ۱ از ۲', first.content.decode())
        self.assertIn('class="tclass-table"', self._get(view='list').content.decode())


class TeacherStudentScopeTests(_ScopedOrgFixture):
    """بخش «دانشجویان» استاد: فقط دانشجویان گروه‌های خودش (مالک یا هم‌تدریس) و فقط آزمون‌های منتشرشده‌ی خودش.

    هر شناسه‌ی دیگر در URL یا فرم ارسالی پاسخ ۴۰۴ می‌گیرد، نه داده‌ی ساختگی یا دانشجوی دیگری.
    """

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        User = get_user_model()
        cls.teacher_user = User.objects.create_user(username='teacher_in_scope', email='t_in@example.com', password='x')
        cls.student2_id = cls._make_profile('student_second', 'دانشجوی دوم')
        _insert('student_profiles', ['user_id', 'student_number'], [cls.student2_id, '40112033'])
        for group_id, student_id in ((cls.group_in_id, cls.student_in_id), (cls.group_in_id, cls.student2_id), (cls.group_out_id, cls.student_out_id)):
            cls._member(group_id, student_id)
        cls.shared_group_id = str(uuid.uuid4())
        _insert('student_groups', ['id', 'teacher_id', 'course_id', 'course_name', 'academic_year'],
                [cls.shared_group_id, cls.teacher_out_id, cls.course_out_id, 'کلاس هم‌تدریس', '1404'])
        _insert('group_teachers', ['group_id', 'teacher_id'], [cls.shared_group_id, cls.teacher_in_id])
        cls.shared_student_id = cls._make_profile('student_shared', 'دانشجوی کلاس هم‌تدریس')
        _insert('student_profiles', ['user_id'], [cls.shared_student_id])
        cls._member(cls.shared_group_id, cls.shared_student_id)

        now = timezone.now()
        day = timedelta(days=1)
        cls.past_exam = cls._exam('آزمون گذشته', cls.teacher_in_id, now - 10 * day, now - 9 * day, cls.group_in_id)
        cls.unfinished_exam = cls._exam('آزمون نیمه‌کاره', cls.teacher_in_id, now - 5 * day, now - 4 * day, cls.group_in_id)
        cls.upcoming_exam = cls._exam('آزمون آینده', cls.teacher_in_id, now + 5 * day, now + 6 * day, cls.group_in_id)
        cls.draft_exam = cls._exam('آزمون پیش‌نویس', cls.teacher_in_id, now - 3 * day, now - 2 * day, cls.group_in_id, published=False)
        cls.other_exam = cls._exam('آزمون استاد دیگر', cls.teacher_out_id, now - 10 * day, now - 9 * day, cls.group_out_id)
        cls._attempt(cls.past_exam, cls.student_in_id, now - 10 * day, 'graded', 15, 20)
        cls._attempt(cls.unfinished_exam, cls.student_in_id, now - 5 * day, 'graded', 12, 20)
        cls._attempt(cls.unfinished_exam, cls.student2_id, now - 5 * day, 'in_progress')
        cls._attempt(cls.other_exam, cls.student_out_id, now - 10 * day, 'graded', 20, 20)

    @staticmethod
    def _member(group_id, student_id):
        _insert('student_group_members', ['id', 'group_id', 'student_user_id', 'full_name', 'national_id', 'student_number'],
                [str(uuid.uuid4()), group_id, student_id, 'عضو', '', ''])

    @staticmethod
    def _exam(title, teacher_id, start_at, end_at, group_id, published=True):
        exam_id = str(uuid.uuid4())
        _insert('exams', ['id', 'title', 'teacher_id', 'start_at', 'end_at', 'is_published', 'is_cancelled'],
                [exam_id, title, teacher_id, start_at, end_at, published, False])
        _insert('exam_assignments', ['id', 'exam_id', 'group_id'], [str(uuid.uuid4()), exam_id, group_id])
        return exam_id

    @staticmethod
    def _attempt(exam_id, student_id, started_at, status, score=None, max_score=None):
        submitted = started_at + timedelta(minutes=30) if status == 'graded' else None
        _insert('exam_attempts', ['id', 'exam_id', 'student_id', 'started_at', 'submitted_at', 'score', 'max_score', 'is_graded', 'status'],
                [str(uuid.uuid4()), exam_id, student_id, started_at, submitted, score, max_score, status == 'graded', status])

    def _get(self, url, data=None):
        self.client.force_login(self.teacher_user)
        return self.client.get(url, data or {})

    def _post(self, url, data):
        self.client.force_login(self.teacher_user)
        return self.client.post(url, data)

    def _notifications(self):
        with connection.cursor() as cursor:
            cursor.execute('SELECT user_id FROM notifications')
            return [str(row[0]) for row in cursor.fetchall()]

    def _notes(self):
        with connection.cursor() as cursor:
            cursor.execute('SELECT id, teacher_id, student_id, body, is_shared FROM teacher_student_notes ORDER BY body')
            return cursor.fetchall()

    def _list_ids(self, **params):
        response = self._get(reverse('core:teacher_students'), params)
        self.assertEqual(response.status_code, 200)
        return {row['id'] for row in response.context['sl']['rows']}

    # -- list
    def test_list_has_own_and_co_taught_students_only(self):
        self.assertEqual(self._list_ids(), {self.student_in_id, self.student2_id, self.shared_student_id})
        self.assertNotContains(self._get(reverse('core:teacher_students')), 'سارا احمدی')

    def test_list_figures_count_only_own_published_exams(self):
        rows = {row['id']: row for row in self._get(reverse('core:teacher_students')).context['sl']['rows']}
        first, second = rows[self.student_in_id], rows[self.student2_id]
        self.assertEqual((first['assigned_count'], first['attended_count'], first['absent_count']), (3, 2, 0))
        self.assertEqual(first['average'], 13.5)
        self.assertEqual(first['last_exam']['title'], 'آزمون نیمه‌کاره')
        self.assertEqual((second['absent_count'], second['unfinished_count']), (1, 1))
        self.assertEqual(second['last_exam']['state'], 'unfinished')

    def test_tabs_and_filters(self):
        self.assertEqual(self._list_ids(tab='absent'), {self.student2_id})
        self.assertEqual(self._list_ids(tab='unfinished'), {self.student2_id})
        self.assertEqual(self._list_ids(tab='review'), {self.student_in_id})  # average 13.5 < 14
        self.assertEqual(self._list_ids(q='40112033'), {self.student2_id})
        self.assertEqual(self._list_ids(group=self.shared_group_id), {self.shared_student_id})
        self.assertEqual(len(self._list_ids(group=self.group_out_id)), 3)  # another teacher's group is ignored

    # -- profile
    def test_profile_tabs_render_for_own_and_co_taught_students(self):
        for student_id in (self.student_in_id, self.shared_student_id):
            for tab in ('overview', 'courses', 'exams', 'analytics', 'notes'):
                response = self._get(reverse('core:teacher_student_profile', args=[student_id]), {'tab': tab})
                self.assertEqual(response.status_code, 200, (student_id, tab))
                self.assertEqual(response.context['profile']['student']['id'], student_id)

    def test_exams_tab_states_and_scope(self):
        url = reverse('core:teacher_student_profile', args=[self.student2_id])
        states = {row['title']: row['state'] for row in self._get(url, {'tab': 'exams'}).context['profile']['exam_rows']}
        self.assertEqual(states, {'آزمون گذشته': 'absent', 'آزمون نیمه‌کاره': 'unfinished', 'آزمون آینده': 'upcoming'})
        filtered = self._get(url, {'tab': 'exams', 'state': 'absent'}).context['profile']['exam_rows']
        self.assertEqual([row['title'] for row in filtered], ['آزمون گذشته'])

    def test_courses_tab_lists_only_own_groups(self):
        response = self._get(reverse('core:teacher_student_profile', args=[self.student_in_id]), {'tab': 'courses'})
        self.assertEqual([group['id'] for group in response.context['profile']['groups']], [self.group_in_id])

    def test_analytics_compares_with_class(self):
        response = self._get(reverse('core:teacher_student_profile', args=[self.student_in_id]), {'tab': 'analytics'})
        analytics = response.context['profile']['analytics']
        self.assertTrue(analytics['has_data'])
        self.assertEqual([point['student'] for point in analytics['points']], [15.0, 12.0])
        self.assertEqual(analytics['points'][0]['minutes'], 30.0)

    def test_other_or_unknown_student_is_404_not_a_stand_in(self):
        for student_id in (self.student_out_id, str(uuid.uuid4()), '14021057', 'sara'):
            for name in ('core:teacher_student_profile', 'core:teacher_student_report'):
                self.assertEqual(self._get(reverse(name, args=[student_id])).status_code, 404, (name, student_id))

    def test_group_pages_redirect_to_filtered_list_or_404(self):
        for name in ('core:teacher_group_students', 'core:teacher_group_activity'):
            response = self._get(reverse(name, args=[self.shared_group_id]))
            self.assertRedirects(response, f"{reverse('core:teacher_students')}?group={self.shared_group_id}")
            self.assertEqual(self._get(reverse(name, args=[self.group_out_id])).status_code, 404)
            self.assertEqual(self._get(reverse(name, args=['nursing-01'])).status_code, 404)

    def test_report_card(self):
        response = self._get(reverse('core:teacher_student_report', args=[self.student_in_id]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'آزمون گذشته')
        self.assertNotContains(response, 'آزمون استاد دیگر')
        self.assertNotContains(response, 'آزمون پیش‌نویس')

    # -- messages
    def test_message_reaches_only_own_students(self):
        url = reverse('core:teacher_students_message')
        response = self._post(url, {'student_ids': [self.student_in_id, self.shared_student_id], 'title': 'یادآوری', 'body': 'سلام'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(sorted(self._notifications()), sorted([self.student_in_id, self.shared_student_id]))
        tampered = self._post(url, {'student_ids': [self.student_in_id, self.student_out_id], 'title': 'x', 'body': 'y'})
        self.assertEqual(tampered.status_code, 404)
        self._post(url, {'student_ids': [self.student_in_id], 'title': '', 'body': 'y'})
        self.assertEqual(len(self._notifications()), 2)

    # -- notes
    def test_notes_are_private_and_scoped(self):
        url = reverse('core:teacher_student_note_create', args=[self.student_in_id])
        self._post(url, {'body': 'الف: نیاز به تمرین بیشتر', 'action': 'save'})
        self._post(url, {'body': 'ب: آفرین', 'action': 'share'})
        self._post(url, {'body': 'x' * 1001, 'action': 'save'})
        notes = self._notes()
        self.assertEqual(len(notes), 2)
        self.assertEqual(self._notifications(), [self.student_in_id])  # only the shared note reaches the student
        self.assertEqual(self._post(reverse('core:teacher_student_note_create', args=[self.student_out_id]), {'body': 'x'}).status_code, 404)

        foreign_note = str(uuid.uuid4())
        _insert('teacher_student_notes', ['id', 'teacher_id', 'student_id', 'body'], [foreign_note, self.teacher_out_id, self.student_in_id, 'ج: خصوصی استاد دیگر'])
        page = self._get(reverse('core:teacher_student_profile', args=[self.student_in_id]), {'tab': 'notes'})
        self.assertNotContains(page, 'خصوصی استاد دیگر')
        self.assertEqual(self._post(reverse('core:teacher_student_note_delete', args=[self.student_in_id, foreign_note]), {}).status_code, 404)
        self._post(reverse('core:teacher_student_note_delete', args=[self.student_in_id, str(notes[0][0])]), {})
        self.assertEqual([row[3] for row in self._notes()], ['ب: آفرین', 'ج: خصوصی استاد دیگر'])

    # -- export
    def test_export_contains_only_filtered_own_students(self):
        response = self._get(reverse('core:teacher_students_export'), {'tab': 'absent'})
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment;', response['Content-Disposition'])
        if response['Content-Disposition'].endswith('.csv"'):
            text = response.content.decode('utf-8-sig')
        else:
            from openpyxl import load_workbook
            sheet = load_workbook(io.BytesIO(response.content)).active
            text = ' '.join(str(cell.value) for row in sheet.iter_rows() for cell in row)
        self.assertIn('دانشجوی دوم', text)
        self.assertNotIn('دانشجوی داخل محدوده', text)
        self.assertNotIn('دانشجوی خارج محدوده', text)

    def test_formula_like_cells_are_neutralised(self):
        from .teacher_students import _cell

        self.assertEqual(_cell('=HYPERLINK("x")'), "'=HYPERLINK(\"x\")")
        self.assertEqual(_cell('علی'), 'علی')


class TeacherSettingsTests(_ScopedOrgFixture):
    """تنظیمات پنل استاد: سه تب، ذخیره فقط فیلدهای همان تب، و ردکردن مقدارهای خارج از گزینه‌ها سمت سرور."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        User = get_user_model()
        cls.teacher_user = User.objects.create_user(username='teacher_in_scope', email='t_in@example.com', password='x')

    def setUp(self):
        self.client.force_login(self.teacher_user)

    def _stored(self):
        from apps.core.views import _teacher_settings_payload
        return _teacher_settings_payload(self.teacher_in_id)

    def _post(self, tab, **data):
        return self.client.post(reverse('core:teacher_settings'), {'settings_tab': tab, **data})

    def test_tabs_render_and_old_tab_names_map_to_new_ones(self):
        for tab, marker in (('account', 'محیط کار'), ('exam', 'پیش‌فرض آزمون'), ('security', 'حریم خصوصی و نمایش اطلاعات')):
            response = self.client.get(reverse('core:teacher_settings'), {'tab': tab})
            self.assertEqual(response.status_code, 200)
            self.assertIn(marker, response.content.decode())
        self.assertEqual(self.client.get(reverse('core:teacher_settings'), {'tab': 'bank'}).context['teacher_settings_tab'], 'exam')
        self.assertEqual(self.client.get(reverse('core:teacher_settings'), {'tab': 'notifications'}).context['teacher_settings_tab'], 'account')
        self.assertEqual(self.client.get(reverse('core:teacher_settings'), {'tab': 'bogus'}).context['teacher_settings_tab'], 'account')

    def test_saves_only_the_posted_tab_and_validates_choices(self):
        self._post('exam', default_exam_duration='90', negative_marking='on', sharing_mode='group', auto_archive='on')
        stored = self._stored()
        self.assertEqual((stored['default_exam_duration'], stored['negative_marking'], stored['sharing_mode'], stored['auto_archive']), ('90', True, 'group', True))
        self.assertFalse(stored['text_feedback'])  # an unchecked switch on the posted tab turns off
        self.assertTrue(stored['quick_help'])  # switches on other tabs keep their value

        self._post('exam', default_exam_duration='9999', sharing_mode='<script>')
        stored = self._stored()
        self.assertEqual((stored['default_exam_duration'], stored['sharing_mode']), ('90', 'group'))

    def test_account_tab_validates_email_and_quiet_hours(self):
        self._post('account', theme='dark', delivery_mode='daily', notification_email='me@uni.ac.ir', quiet_start='23:30', quiet_end='06:15')
        stored = self._stored()
        self.assertEqual(
            (stored['theme'], stored['delivery_mode'], stored['notification_email'], stored['quiet_start'], stored['quiet_end']),
            ('dark', 'daily', 'me@uni.ac.ir', '23:30', '06:15'),
        )
        response = self._post('account', notification_email='not-an-email', quiet_start='25:99')
        stored = self._stored()
        self.assertEqual((stored['notification_email'], stored['quiet_start']), ('me@uni.ac.ir', '23:30'))
        self.assertIn('ایمیل اعلان‌ها معتبر نیست', self.client.get(response['Location']).content.decode())

    def test_security_tab_shows_no_fake_sessions(self):
        page = self.client.get(reverse('core:teacher_settings'), {'tab': 'security'}).content.decode()
        for fake in ('Windows 11', 'Android · Chrome', 'حساب شما ایمن است'):
            self.assertNotIn(fake, page)
        self.assertIn('در حال توسعه', page)


class SuperAdminExamEditScopeTests(_ScopedOrgFixture):
    """صفحه‌ی ویرایش آزمون مدیر آموزشی: فهرست‌ها و هر شناسه‌ی ارسالی باید داخل محدوده‌ی او باشد."""

    def _post(self, step, **data):
        self.client.force_login(self.manager_user)
        return self.client.post(reverse('core:super_admin_exam_edit', args=[self.exam_in_id]), {'edit_step': step, **data})

    def test_pickers_are_scoped(self):
        self.client.force_login(self.manager_user)
        ctx = self.client.get(reverse('core:super_admin_exam_edit', args=[self.exam_in_id])).context
        self.assertNotIn(self.course_out_id, {c['id'] for c in ctx['courses']})
        self.assertNotIn(self.teacher_out_id, {t['id'] for t in ctx['teachers']})
        self.assertNotIn(self.group_out_id, {g['id'] for g in ctx['groups']})
        self.assertNotIn(self.student_out_id, {s['id'] for s in ctx['students']})
        self.assertNotIn(self.question_out_id, {q['id'] for q in ctx['questions']})
        self.assertIn(self.group_in_id, {g['id'] for g in ctx['groups']})

    def test_out_of_scope_ids_are_rejected_per_step(self):
        cases = (
            ('1', {'title': 'x', 'course_id': self.course_out_id}),
            ('1', {'title': 'x', 'teacher_id': self.teacher_out_id}),
            ('2', {'group_ids': [self.group_out_id]}),
            ('2', {'student_ids': [self.student_out_id]}),
            ('3', {'question_ids': [self.question_out_id]}),
        )
        for step, data in cases:
            self.assertEqual(self._post(step, **data).status_code, 403, data)
        with connection.cursor() as cursor:
            cursor.execute('SELECT course_id, teacher_id FROM exams WHERE id = %s', [self.exam_in_id])
            self.assertEqual(cursor.fetchone(), (self.course_in_id, self.teacher_in_id))
            cursor.execute('SELECT COUNT(*) FROM exam_assignments WHERE exam_id = %s', [self.exam_in_id])
            self.assertEqual(cursor.fetchone()[0], 0)

    def test_in_scope_assignment_is_saved(self):
        response = self._post('2', group_ids=[self.group_in_id], student_ids=[self.student_in_id])
        self.assertEqual(response.status_code, 302)
        with connection.cursor() as cursor:
            cursor.execute('SELECT COUNT(*) FROM exam_assignments WHERE exam_id = %s', [self.exam_in_id])
            self.assertEqual(cursor.fetchone()[0], 2)


class TeacherExamReportTests(_ScopedOrgFixture):
    """صفحه‌های «نتایج» و «پایش زنده» استاد: داده‌ی واقعی آزمون خود او، و ۴۰۴ برای آزمون استاد دیگر."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        User = get_user_model()
        cls.teacher_user = User.objects.create_user(username='teacher_in_scope', email='t_in@example.com', password='x')

    def _attempt(self, exam_id, status, score=None, max_score=20):
        _insert('exam_attempts', ['id', 'exam_id', 'student_id', 'started_at', 'status', 'score', 'max_score', 'is_graded'],
                [str(uuid.uuid4()), exam_id, self.student_in_id, timezone.now(), status, score, max_score, score is not None])

    def test_results_use_real_scores(self):
        self._attempt(self.exam_in_id, 'graded', 18)
        self._attempt(self.exam_in_id, 'graded', 6)
        self._attempt(self.exam_in_id, 'submitted')
        self.client.force_login(self.teacher_user)
        response = self.client.get(reverse('core:teacher_results'), {'exam': self.exam_in_id})
        report = response.context['report']
        self.assertEqual((report['average'], report['highest'], report['lowest'], report['ungraded']), (12.0, 18.0, 6.0, 1))
        self.assertEqual(report['pass_rate'], 50)
        self.assertEqual([bucket['count'] for bucket in report['buckets']], [0, 1, 0, 0, 1])
        page = response.content.decode()
        for leftover in ('۶۴.۲', '۹۸.۵', '۷۲.۶٪', 'انتشار نتایج برای دانشجویان'):
            self.assertNotIn(leftover, page)

    def test_results_default_to_exam_with_submissions(self):
        self._attempt(self.exam_in_id, 'graded', 15)
        self.client.force_login(self.teacher_user)
        response = self.client.get(reverse('core:teacher_results'))
        self.assertEqual(str(response.context['report_exam']['id']), self.exam_in_id)

    def test_monitoring_counts_live_attempts(self):
        from datetime import timedelta
        with connection.cursor() as cursor:
            cursor.execute('UPDATE exams SET start_at = %s, end_at = %s WHERE id = %s',
                           [timezone.now() - timedelta(minutes=10), timezone.now() + timedelta(minutes=50), self.exam_in_id])
        self._attempt(self.exam_in_id, 'in_progress')
        self.client.force_login(self.teacher_user)
        response = self.client.get(reverse('core:teacher_monitoring'))
        report = response.context['report']
        self.assertTrue(report['is_running'])
        self.assertEqual((report['in_progress'], report['submitted'], len(report['rows'])), (1, 0, 1))
        self.assertGreater(report['remaining_seconds'], 0)
        page = response.content.decode()
        for leftover in ('اصول حسابداری', 'قفل کردن آزمون', 'سارا محمدی قطع اتصال داشت'):
            self.assertNotIn(leftover, page)

    def test_other_teachers_exam_is_404(self):
        self.client.force_login(self.teacher_user)
        for name in ('core:teacher_results', 'core:teacher_monitoring'):
            self.assertEqual(self.client.get(reverse(name), {'exam': self.exam_out_id}).status_code, 404)


class StudentExamAccessFallbackTests(_ScopedOrgFixture):
    """دانشجو با دانستن شناسه نباید به آزمون دیگران برسد؛ قبلاً هر دو تابع دسترسی بدون شرط به همه‌ی آزمون‌ها برمی‌گشتند."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        User = get_user_model()
        cls.student_user = User.objects.create_user(username='student_in_scope', email='s_in@example.com', password='x')
        with connection.cursor() as cursor:
            cursor.execute('UPDATE exams SET is_published = %s, lifecycle_status = %s WHERE id = %s', [True, 'published', cls.exam_out_id])
        _insert('exam_assignments', ['id', 'exam_id', 'group_id'], [str(uuid.uuid4()), cls.exam_out_id, cls.group_out_id])

    def test_other_groups_exam_is_not_listed_or_reachable(self):
        from apps.core.views import _sx_accessible_exams
        self.assertEqual(_sx_accessible_exams(self.student_in_id), [])
        self.client.force_login(self.student_user)
        for name in ('core:student_exam_detail', 'core:student_exam_entry', 'core:student_exam_start'):
            self.assertEqual(self.client.get(reverse(name, args=[self.exam_out_id])).status_code, 404, name)
        with connection.cursor() as cursor:
            cursor.execute('SELECT COUNT(*) FROM exam_attempts WHERE exam_id = %s', [self.exam_out_id])
            self.assertEqual(cursor.fetchone()[0], 0)

    def test_student_keeps_access_to_an_exam_they_already_attempted(self):
        from apps.core.views import _sx_exam_row
        _insert('exam_attempts', ['id', 'exam_id', 'student_id', 'started_at', 'status'],
                [str(uuid.uuid4()), self.exam_out_id, self.student_in_id, timezone.now(), 'submitted'])
        self.assertIsNotNone(_sx_exam_row(self.exam_out_id, self.student_in_id))
        self.assertIsNone(_sx_exam_row(self.exam_out_id, self.student_out_id))


class UnassignedExamAccessTests(_ScopedOrgFixture):
    """آزمونی که هیچ تخصیصی ندارد برای هیچ دانشجویی باز نیست؛ مسیرهای ایجاد سریع گروه را به‌عنوان مخاطب ذخیره می‌کنند."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        User = get_user_model()
        cls.student_user = User.objects.create_user(username='student_in_scope', email='s_in@example.com', password='x')
        with connection.cursor() as cursor:
            cursor.execute('UPDATE exams SET is_published = %s, lifecycle_status = %s WHERE id = %s', [True, 'published', cls.exam_in_id])

    def _accessible(self):
        from apps.core.views import _sx_accessible_exams
        return {str(row['id']) for row in _sx_accessible_exams(self.student_in_id)}

    def test_published_exam_without_assignment_is_hidden(self):
        self.assertNotIn(self.exam_in_id, self._accessible())
        self.client.force_login(self.student_user)
        self.assertEqual(self.client.get(reverse('core:student_exam_detail', args=[self.exam_in_id])).status_code, 404)

    def test_group_assignment_opens_it_for_members_only(self):
        _insert('student_group_members', ['id', 'group_id', 'student_user_id', 'full_name', 'national_id', 'student_number'],
                [str(uuid.uuid4()), self.group_in_id, self.student_in_id, 'دانشجو', '1', '1'])
        _insert('exam_assignments', ['id', 'exam_id', 'group_id'], [str(uuid.uuid4()), self.exam_in_id, self.group_in_id])
        self.assertIn(self.exam_in_id, self._accessible())
        from apps.core.views import _sx_accessible_exams
        self.assertNotIn(self.exam_in_id, {str(row['id']) for row in _sx_accessible_exams(self.student_out_id)})

    def test_quick_create_saves_the_group_as_audience(self):
        self.client.force_login(self.manager_user)
        response = self.client.post(reverse('core:super_admin_exams'), {
            'exam_action': 'create', 'group_id': self.group_in_id, 'title': 'آزمون سریع',
            'question_ids': [self.question_in_id], 'duration_minutes': '30',
        })
        exam_id = response.json().get('exam_id')
        self.assertTrue(response.json()['ok'], response.json())
        with connection.cursor() as cursor:
            cursor.execute('SELECT group_id FROM exam_assignments WHERE exam_id = %s', [exam_id])
            self.assertEqual([row[0] for row in cursor.fetchall()], [self.group_in_id])


class TeacherExamEditTests(_ScopedOrgFixture):
    """ویرایش آزمون استاد با همان ویزارد: فقط آزمون خودش، فقط تا پیش از انتشار و شرکت دانشجو."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        User = get_user_model()
        cls.teacher_user = User.objects.create_user(username='teacher_in_scope', email='t_in@example.com', password='x')

    def setUp(self):
        self.client.force_login(self.teacher_user)
        self.client.post(reverse('core:teacher_exam_create'), {
            'title': 'آزمون قابل ویرایش', 'course_id': self.course_in_id, 'group_ids': [self.group_in_id],
            'question_ids': [self.question_in_id, self.question_in2_id], 'duration_minutes': '60',
            'start_at': '2026-11-10T09:00', 'tags': 'قلب، عروق', 'publish_choice': 'submit',
        })
        with connection.cursor() as cursor:
            cursor.execute("SELECT id FROM exams WHERE title = 'آزمون قابل ویرایش'")
            self.exam_id = cursor.fetchone()[0]
        self.url = reverse('core:teacher_exam_edit', args=[self.exam_id])

    def _rows(self, sql):
        with connection.cursor() as cursor:
            cursor.execute(sql, [self.exam_id])
            return cursor.fetchall()

    def test_get_prefills_wizard_with_saved_exam(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        initial = response.context['wizard_initial']
        self.assertEqual(initial['fields']['title'], 'آزمون قابل ویرایش')
        self.assertEqual(initial['fields']['start_at'], '2026-11-10T09:00')
        self.assertEqual(initial['question_ids'], [self.question_in_id, self.question_in2_id])
        self.assertEqual(initial['group_ids'], [self.group_in_id])
        self.assertEqual(initial['tags'], ['قلب', 'عروق'])
        self.assertEqual(initial['publish_choice'], 'submit')
        self.assertIn('id="wizard-initial"', response.content.decode())
        self.assertIn('ویرایش آزمون', response.content.decode())

    def test_post_updates_in_place_and_replaces_children(self):
        response = self.client.post(self.url, {
            'title': 'آزمون ویرایش‌شده', 'course_id': self.course_in_id, 'group_ids': [self.group_in_id],
            'question_ids': [self.question_in2_id], 'duration_minutes': '45', 'publish_choice': 'draft',
        })
        self.assertRedirects(response, reverse('core:teacher_exams'), fetch_redirect_response=False)
        self.assertEqual(self._rows('SELECT title, duration_minutes, lifecycle_status, teacher_id FROM exams WHERE id = %s'),
                         [('آزمون ویرایش‌شده', 45, 'draft', self.teacher_in_id)])
        self.assertEqual(self._rows('SELECT question_id FROM exam_questions WHERE exam_id = %s'), [(self.question_in2_id,)])
        self.assertEqual(self._rows('SELECT group_id FROM exam_assignments WHERE exam_id = %s'), [(self.group_in_id,)])
        with connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) FROM exams WHERE title LIKE 'آزمون%%ویرایش%%'")
            self.assertEqual(cursor.fetchone()[0], 1)

    def test_out_of_scope_ids_rejected_on_edit(self):
        response = self.client.post(self.url, {
            'title': 'x', 'course_id': self.course_in_id, 'group_ids': [self.group_out_id], 'question_ids': [self.question_in_id],
        })
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self._rows('SELECT title FROM exams WHERE id = %s'), [('آزمون قابل ویرایش',)])

    def test_published_or_started_exam_is_not_editable(self):
        _insert('exam_attempts', ['id', 'exam_id', 'student_id', 'started_at', 'status'],
                [str(uuid.uuid4()), self.exam_id, self.student_in_id, timezone.now(), 'in_progress'])
        self.assertRedirects(self.client.get(self.url), reverse('core:teacher_exams'), fetch_redirect_response=False)
        self.assertRedirects(self.client.post(self.url, {'title': 'x'}), reverse('core:teacher_exams'), fetch_redirect_response=False)
        with connection.cursor() as cursor:
            cursor.execute('UPDATE exams SET is_published = %s, lifecycle_status = %s WHERE id = %s', [True, 'published', self.exam_out_id])
        self.assertEqual(self.client.get(reverse('core:teacher_exam_edit', args=[self.exam_out_id])).status_code, 404)

    def test_exam_list_links_to_edit_only_when_editable(self):
        with connection.cursor() as cursor:
            cursor.execute('UPDATE exams SET is_published = %s, lifecycle_status = %s WHERE id = %s', [True, 'published', self.exam_in_id])
        page = self.client.get(reverse('core:teacher_exams')).content.decode()
        self.assertIn(self.url, page)
        self.assertNotIn(reverse('core:teacher_exam_edit', args=[self.exam_in_id]), page)

    def test_exam_with_items_outside_teacher_scope_is_not_editable(self):
        _insert('exam_questions', ['id', 'exam_id', 'question_id', 'points', 'order_index'],
                [str(uuid.uuid4()), self.exam_id, self.question_out_id, 1, 9])
        response = self.client.get(self.url)
        self.assertRedirects(response, reverse('core:teacher_exams'), fetch_redirect_response=False)
        self.assertEqual(len(self._rows('SELECT question_id FROM exam_questions WHERE exam_id = %s')), 3)  # nothing dropped
